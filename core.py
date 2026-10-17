"""Portable, project-scoped data store and Excel interchange for VoiceDrama Studio."""
from __future__ import annotations
import copy
import json
import os
import re
import shutil
import threading
import uuid
from datetime import datetime
from pathlib import Path
from difflib import SequenceMatcher

import openpyxl

ROOT = Path(__file__).resolve().parent
WORKSPACE = Path(os.getenv("VDS_WORKSPACE", str(ROOT / "workspace"))).resolve()
LOCK = threading.RLock()


# TTS installation paths are PC-wide, independent from script/project state.
GLOBAL_TTS_DEFAULTS = {
    "irodori_dir": "D:\\Irodori-TTS",
    "irodori_python": "",
    "qwen_python": "",
}
GLOBAL_TTS_KEYS = tuple(GLOBAL_TTS_DEFAULTS)


def global_settings_path():
    return WORKSPACE / "global_settings.json"


def _atomic_save_global(data):
    path = global_settings_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temp, path)


def load_global_settings():
    """Load common TTS paths. Import old project-specific settings only once.

    Prefer the newest non-default configured value for each key; the original
    project files are never mutated by migration. Once global_settings.json
    exists, project files can no longer overwrite global settings.
    """
    with LOCK:
        path = global_settings_path()
        if path.exists():
            data = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(data, dict) or any(k not in data for k in GLOBAL_TTS_KEYS):
                raise ValueError("共通環境設定が破損しています。global_settings.jsonを確認してください")
            return data
        data = {**GLOBAL_TTS_DEFAULTS, "revision": 1}
        chosen = set()
        for project in list_projects():
            try:
                legacy = load_project(project["id"]).get("settings", {})
                if not isinstance(legacy, dict):
                    continue
                for key in GLOBAL_TTS_KEYS:
                    value = legacy.get(key)
                    if (key not in chosen and isinstance(value, str) and value.strip()
                            and value.strip() != GLOBAL_TTS_DEFAULTS[key]):
                        data[key] = value.strip()
                        chosen.add(key)
            except (OSError, ValueError, KeyError):
                continue
        _atomic_save_global(data)
        return data


def save_global_settings(settings, expected_revision=None):
    """Update shared paths with optimistic revision check, without modifying projects."""
    with LOCK:
        old = load_global_settings()
        if expected_revision is not None and int(expected_revision) != old["revision"]:
            raise RuntimeError("共通環境設定が別の画面で変更されました。再読み込みしてください。")
        new = {}
        for key in GLOBAL_TTS_KEYS:
            value = settings.get(key)
            if not isinstance(value, str) or len(value) > 2048 or any(ord(ch) < 32 for ch in value):
                raise ValueError(f"{key} の設定が正しくありません")
            new[key] = value.strip()
        if all(old[key] == new[key] for key in GLOBAL_TTS_KEYS):
            return old
        result = {**new, "revision": old["revision"] + 1}
        _atomic_save_global(result)
        return result


def uid(prefix="id"):
    return prefix + "_" + uuid.uuid4().hex[:16]


def now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def workdir(project_id):
    if not re.fullmatch(r"p_[a-f0-9]{16}", str(project_id)):
        raise ValueError("Invalid project ID")
    return WORKSPACE / project_id


def asset_subdir(kind):
    """Human-readable physical folders; asset IDs/metadata stay unchanged."""
    return "reference_audio" if kind == "reference" else "dialogue_audio" if kind == "voice" else "assets"


def asset_path(project_id, asset):
    """Canonical location for a registered asset, never a user-provided path."""
    filename = asset.get("filename", "")
    if not isinstance(filename, str) or not re.fullmatch(
        r"asset_[a-f0-9]{16}\.(?:wav|mp3|flac|ogg|m4a|aac|opus)", filename
    ):
        raise ValueError("安全でない音声ファイル名です")
    return workdir(project_id) / asset_subdir(asset.get("kind")) / filename


def _migrate_audio_folders(project_id, state):
    """Move old reference/dialogue files once, without changing project JSON.

    Only actual asset entries are moved. An interrupted migration is restartable:
    files already moved remain at their destination, missing files remain missing.
    On an error, completed moves in this pass are rolled back.
    """
    root = workdir(project_id)
    moves = []
    try:
        for asset in state.get("assets", []):
            if asset.get("kind") not in ("reference", "voice"):
                continue
            new = asset_path(project_id, asset)
            legacy = root / "assets" / asset["filename"]
            if new.is_file() or not legacy.is_file():
                continue
            new.parent.mkdir(parents=True, exist_ok=True)
            if new.exists():
                raise RuntimeError(f"音声保存先が使用されています: {new.name}")
            # Both paths are under the same project root, allowing atomic rename.
            os.replace(legacy, new)
            moves.append((legacy, new))
    except Exception:
        for original, moved in reversed(moves):
            if moved.is_file() and not original.exists():
                os.replace(moved, original)
        raise


def state_path(project_id):
    return workdir(project_id) / "project.json"


def default_state(title, project_id):
    scene = uid('scene')
    return {
        "id": project_id, "title": title.strip() or "無題の作品", "updatedAt": now(), "revision": 0,
        "scenes": [{"id": scene, "name": "シーン1"}], "lines": [],
        "tracks": [
            {"id":"dialogue", "name":"台詞", "kind":"dialogue"},
            {"id":"sfx", "name":"効果音", "kind":"sfx"},
            {"id":"bgm", "name":"BGM", "kind":"bgm"},
            {"id":"ambience", "name":"環境音", "kind":"ambience"},
        ],
        "clips": [], "assets": [], "voices": [],
        "settings": {"defaultGap":0.65},
    }


def list_projects():
    WORKSPACE.mkdir(parents=True, exist_ok=True)
    out=[]
    for path in WORKSPACE.glob("p_*/project.json"):
        try:
            s=json.loads(path.read_text(encoding="utf-8"))
            out.append({"id":s["id"], "title":s.get("title", ""), "updatedAt":s.get("updatedAt", "")})
        except (OSError,ValueError,KeyError):
            pass
    return sorted(out,key=lambda p:p["updatedAt"],reverse=True)


def create_project(title):
    with LOCK:
        project_id=uid('p')
        state=default_state(title,project_id)
        workdir(project_id).mkdir(parents=True)
        for folder in ("assets", "dialogue_audio", "reference_audio"):
            (workdir(project_id)/folder).mkdir()
        state=save_project(project_id,state)
        return state


def delete_cast(project_id, name, expected_revision):
    """Remove an unused cast profile; never delete media or rewrite dialogue."""
    if not isinstance(name, str) or not name.strip():
        raise ValueError("削除するキャスト名を指定してください")
    if expected_revision is None:
        raise ValueError("保存リビジョンを指定してください")
    with LOCK:
        state = load_project(project_id)
        if int(state["revision"]) != int(expected_revision):
            raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
        matches = [voice for voice in state["voices"] if voice.get("name") == name]
        if not matches:
            raise FileNotFoundError("指定したキャストが見つかりません")
        if len(matches) != 1:
            raise ValueError("同名のキャストが複数あるため削除できません")
        used_count = sum(line.get("speaker") == name for line in state["lines"])
        if used_count:
            raise RuntimeError(
                f"「{name}」は{used_count}件の台詞で使用中です。"
                "台詞の話者を変更してから削除してください。")
        # Do not modify assets, takes, clips, or other projects. Reference audio
        # remains in the library, even if this was its last cast reference.
        state["voices"] = [v for v in state["voices"] if v.get("name") != name]
        return save_project(project_id, state, expected_revision=expected_revision)


def load_project(project_id):
    with LOCK:
        p=state_path(project_id)
        if not p.is_file():
            raise FileNotFoundError("Project not found")
        state=json.loads(p.read_text(encoding="utf-8"))
        _migrate_audio_folders(project_id, state)
        return state


def save_project(project_id, state, expected_revision=None, *, allow_new_assets=False):
    with LOCK:
        path=state_path(project_id)
        if str(state.get("id"))!=project_id:
            raise ValueError("Project ID mismatch")
        previous=json.loads(path.read_text(encoding="utf-8")) if path.is_file() else None
        if previous and expected_revision is not None and int(previous["revision"]) != int(expected_revision):
            raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
        state=copy.deepcopy(state)
        validate_state(state, previous, allow_new_assets=allow_new_assets)
        state["revision"]=(int(previous.get("revision",0))+1) if previous else 1
        state["updatedAt"]=now()
        path.parent.mkdir(parents=True,exist_ok=True)
        temp=path.with_suffix(".tmp")
        temp.write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding="utf-8")
        os.replace(temp,path)
        return state


def export_scene_selected_takes_zip(project_id, scene_id, destination):
    """Archive exactly the current scene's selected takes in script order.

    The persistent file path and asset ID are never used as ZIP entry names.
    Explicit per-line filenames take priority; duplicate names get numbered.
    No project data is modified.
    """
    import unicodedata
    from zipfile import ZipFile, ZIP_DEFLATED

    def safe_name(value, extension, fallback):
        name = str(value or "").replace("\\", "/").split("/")[-1].strip()
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f\x7f]', "_", name).strip(" .")
        if not name:
            name = fallback + extension
        stem = name[:-len(extension)] if name.lower().endswith(extension.lower()) else Path(name).stem
        if not stem or stem.upper().split(".")[0] in {
                "CON", "PRN", "AUX", "NUL",
                *[f"COM{i}" for i in range(1, 10)],
                *[f"LPT{i}" for i in range(1, 10)]}:
            stem = fallback
        return stem[:160].rstrip(" .") + extension

    with LOCK:
        state = load_project(project_id)
        if scene_id not in {scene["id"] for scene in state["scenes"]}:
            raise ValueError("指定したシーンが見つかりません")
        assets = {asset["id"]: asset for asset in state["assets"]}
        items = []
        used_names = set()
        for line in state["lines"]:
            if line.get("sceneId") != scene_id or not line.get("selectedTakeId"):
                continue
            take = next((t for t in line.get("takes", [])
                         if t.get("id") == line["selectedTakeId"]), None)
            if take is None:
                continue
            asset = assets.get(take.get("assetId"))
            if asset is None:
                raise ValueError("選択テイクに対応する音声素材がありません")
            filename = asset.get("filename", "")
            if not isinstance(filename, str) or Path(filename).name != filename:
                raise ValueError("保存されている音声のパスが不正です")
            audio_path = asset_path(project_id, asset)
            if not audio_path.is_file():
                raise FileNotFoundError(f"音声ファイルが見つかりません: {line['id']}")
            extension = audio_path.suffix.lower()
            if extension not in (".wav", ".mp3", ".flac", ".ogg", ".m4a", ".aac", ".opus"):
                raise ValueError("音声形式が正しくありません")
            preferred = line.get("outputFilename") or asset.get("name") or filename
            name = safe_name(preferred, extension, line["id"])
            stem = name[:-len(extension)]
            index = 2
            while unicodedata.normalize("NFC", name).casefold() in used_names:
                name = f"{stem} ({index}){extension}"
                index += 1
            used_names.add(unicodedata.normalize("NFC", name).casefold())
            items.append((audio_path, name))
        if not items:
            raise ValueError("このシーンに選択済みの音声テイクがありません")
        with ZipFile(destination, "w", compression=ZIP_DEFLATED, compresslevel=3,
                     allowZip64=True) as archive:
            for physical, public_name in items:
                archive.write(physical, arcname=public_name)
        return len(items)


def generated_audio_filename(line, engine="irodori"):
    """Name the user-facing WAV; the file stored on disk retains its opaque asset ID."""
    requested = line.get("outputFilename", "") or ""
    if not isinstance(requested, str):
        raise ValueError("音声ファイル名は文字列で指定してください")
    requested = requested.strip()
    if not requested:
        return f'{line["id"]} - {engine}.wav'
    if len(requested) > 120:
        raise ValueError("音声ファイル名は120文字以内で指定してください")
    if (any(ch in requested for ch in '/\\<>:\"|?*')
            or any(ord(ch) < 32 or ord(ch) == 127 for ch in requested)
            or requested.startswith('.') or requested.endswith(('.', ' '))):
        raise ValueError("ファイル名に使用できない文字が含まれています")
    stem = requested[:-4] if requested.lower().endswith('.wav') else requested
    if not stem or '.' in stem or stem.endswith((' ', '.')):
        raise ValueError("拡張子は.wavのみ指定できます")
    if stem.upper().split('.')[0] in {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)], *[f"LPT{i}" for i in range(1, 10)]}:
        raise ValueError("Windowsで使用できない予約名です")
    return stem + ".wav"


def validate_state(s, prev, *, allow_new_assets=False):
    for k in ("scenes","lines","clips","tracks","assets","voices"):
        if not isinstance(s.get(k),list):
            raise ValueError(f"{k} must be a list")
    for k in ("lines","scenes","clips","assets"):
        vals=[v.get("id") for v in s[k]]
        if len(vals)!=len(set(vals)) or not all(vals):
            raise ValueError(f"Duplicate or missing IDs in {k}")
    scene_ids={v["id"] for v in s["scenes"]}
    asset_ids={v["id"] for v in s["assets"]}
    for l in s["lines"]:
        if l.get("sceneId") not in scene_ids: raise ValueError("Line scene mismatch")
        if len(str(l.get("text","")))>10000: raise ValueError("Line too long")
        generated_audio_filename(l)  # Validate optional per-dialogue WAV name.
        for take in l.get("takes",[]):
            if take.get("assetId") not in asset_ids: raise ValueError("Take asset missing")
    track_ids={x["id"] for x in s["tracks"]}
    for c in s["clips"]:
        if c.get("sceneId") not in scene_ids or c.get("trackId") not in track_ids:
            raise ValueError("Clip track/scene mismatch")
        if c.get("assetId") not in asset_ids: raise ValueError("Clip asset missing")
    # New cast references must point to reference recordings. Existing legacy
    # assignments remain readable, and users can uncheck them in the UI.
    asset_by_id = {a["id"]: a for a in s["assets"]}
    legacy_refs = {aid for v in (prev or {}).get("voices", []) for aid in v.get("refs", [])}
    for voice in s["voices"]:
        refs = voice.get("refs", [])
        if len(refs) != len(set(refs)):
            raise ValueError("同じ参照音声が重複しています")
        for aid in refs:
            if aid not in asset_by_id:
                raise ValueError("存在しない参照音声が指定されています")
            if asset_by_id[aid].get("kind") != "reference" and aid not in legacy_refs:
                raise ValueError("台詞音声を参照音声へ登録できません。参照音声として再取り込みしてください")
    # Asset filename is server-controlled: clients cannot replace or invent disk paths.
    if prev:
        originals={a["id"]:a for a in prev["assets"]}
        for a in s["assets"]:
            if (a["id"] not in originals and not allow_new_assets) or (a["id"] in originals and (a.get("filename") != originals[a["id"]].get("filename") or a.get("kind") != originals[a["id"]].get("kind"))):
                raise ValueError("Assets can be added only through the upload API")


# Reference recordings and spoken dialogue are distinct asset roles.
def asset_usage(state, asset_id):
    """Return references to an asset without mutating any project data."""
    uses = {"casts": [], "lines": [], "clips": []}
    for v in state["voices"]:
        if asset_id in v.get("refs", []):
            uses["casts"].append(v.get("name", ""))
    for line in state["lines"]:
        if any(t.get("assetId") == asset_id for t in line.get("takes", [])):
            uses["lines"].append({"id": line["id"], "speaker": line.get("speaker", ""),
                                  "text": line.get("text", "")[:60]})
    for clip in state["clips"]:
        if clip.get("assetId") == asset_id:
            uses["clips"].append(clip["id"])
    return uses


def asset_trash_path(project_id):
    return workdir(project_id) / "trash" / "index.json"


def list_trash(project_id):
    # Validate project even if its trash directory does not exist.
    load_project(project_id)
    path = asset_trash_path(project_id)
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else []


def _write_trash(project_id, entries):
    path = asset_trash_path(project_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, path)


def trash_asset(project_id, asset_id, *, detach=False, expected_revision=None):
    """Move an asset to recoverable trash, optionally detaching every project reference.

    Linked assets are rejected by default. The audio file is moved rather than
    destroyed, and no existing project is silently altered beyond this asset.
    """
    with LOCK:
        state = load_project(project_id)
        if expected_revision is not None and state["revision"] != int(expected_revision):
            raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
        asset = next((a for a in state["assets"] if a["id"] == asset_id), None)
        if asset is None:
            raise FileNotFoundError("音声素材が見つかりません")
        usage = asset_usage(state, asset_id)
        if any(usage.values()) and not detach:
            raise RuntimeError("使用中の音声です。関連付けを確認してください")
        filename = str(asset.get("filename", ""))
        if not re.fullmatch(r"asset_[a-f0-9]{16}\.(wav|mp3|flac|ogg|m4a|aac|opus)", filename):
            raise ValueError("安全でない音声ファイル名です")
        original = asset_path(project_id, asset)
        destination = workdir(project_id) / "trash" / filename
        if destination.exists():
            raise RuntimeError("ゴミ箱内に同名の音声が存在します")
        if not original.is_file():
            raise FileNotFoundError("音声の実体ファイルが見つかりません")
        old_state = copy.deepcopy(state)
        entries = list_trash(project_id)
        state["assets"] = [a for a in state["assets"] if a["id"] != asset_id]
        affected_scenes = set()
        for voice in state["voices"]:
            voice["refs"] = [ref for ref in voice.get("refs", []) if ref != asset_id]
        for line in state["lines"]:
            old_takes = line.get("takes", [])
            kept = [t for t in old_takes if t.get("assetId") != asset_id]
            if len(kept) == len(old_takes):
                continue
            line["takes"] = kept
            affected_scenes.add(line["sceneId"])
            if line.get("selectedTakeId") not in {t["id"] for t in kept}:
                line["selectedTakeId"] = kept[-1]["id"] if kept else None
            dialogue = next((c for c in state["clips"] if c.get("lineId") == line["id"] and c.get("trackId") == "dialogue"), None)
            if dialogue:
                chosen = next((t for t in kept if t["id"] == line["selectedTakeId"]), None)
                fallback = next((a for a in state["assets"] if chosen and a["id"] == chosen["assetId"]), None)
                if fallback:
                    dialogue["assetId"] = fallback["id"]
                    dialogue["duration"] = fallback["duration"]
                else:
                    state["clips"].remove(dialogue)
        state["clips"] = [c for c in state["clips"] if c.get("assetId") != asset_id]
        for sid in affected_scenes:
            reflow_dialogue(state, sid)
        refresh_staleness(state)
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(original), str(destination))
        try:
            result = save_project(project_id, state, expected_revision=expected_revision)
            entries.append({"asset": asset, "trashedAt": now(), "hadReferences": any(usage.values())})
            _write_trash(project_id, entries)
        except Exception:
            # Do not leave project metadata referring to an absent audio file.
            shutil.move(str(destination), str(original))
            if state_path(project_id).exists():
                # If save succeeded but trash index failed, restore old metadata.
                save_project(project_id, old_state, allow_new_assets=True)
            raise
        return {"project": result, "usage": usage, "trashed": asset}


def restore_asset(project_id, asset_id):
    """Restore the file to the library; old links are deliberately NOT recreated."""
    with LOCK:
        state = load_project(project_id)
        entries = list_trash(project_id)
        item = next((x for x in entries if x["asset"]["id"] == asset_id), None)
        if item is None:
            raise FileNotFoundError("ゴミ箱に音声が見つかりません")
        asset = item["asset"]
        if any(x["id"] == asset_id for x in state["assets"]):
            raise RuntimeError("音声IDが既に使用されています")
        filename = str(asset["filename"])
        if not re.fullmatch(r"asset_[a-f0-9]{16}\.(wav|mp3|flac|ogg|m4a|aac|opus)", filename):
            raise ValueError("安全でない音声ファイル名です")
        source = workdir(project_id) / "trash" / filename
        target = asset_path(project_id, asset)
        if not source.is_file() or target.exists():
            raise RuntimeError("復元する音声が見つからないか、保存先に既に存在します")
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(source), str(target))
        state["assets"].append(asset)
        try:
            result = save_project(project_id, state, allow_new_assets=True)
            _write_trash(project_id, [x for x in entries if x is not item])
        except Exception:
            shutil.move(str(target), str(source))
            raise
        return {"project": result, "asset": asset}


def copy_casts(source_project_id, target_project_id, cast_names, *, conflict="skip",
               expected_source_revision=None, expected_target_revision=None):
    """Copy selected casts and their reference WAV/assets into another project.

    Source content is read-only. Each reference asset gets a destination-local ID,
    so future deletes/renames in either project cannot break the other's voices.
    Duplicate reference IDs shared by selected casts are copied only once.
    """
    if source_project_id == target_project_id:
        raise ValueError("コピー元とコピー先には別のプロジェクトを選んでください")
    if conflict not in ("skip", "rename"):
        raise ValueError("同名キャラクターの処理が不正です")
    if not isinstance(cast_names, list) or not 1 <= len(cast_names) <= 200:
        raise ValueError("1〜200件のキャラクターを選んでください")
    if any(not isinstance(n, str) or not n.strip() for n in cast_names) or len(set(cast_names)) != len(cast_names):
        raise ValueError("選択されたキャラクター名が不正または重複しています")

    with LOCK:
        source = load_project(source_project_id)
        target = load_project(target_project_id)
        if (expected_source_revision is not None and
                int(source["revision"]) != int(expected_source_revision)):
            raise RuntimeError("コピー元が更新されています。画面を再読み込みしてください")
        if (expected_target_revision is not None and
                int(target["revision"]) != int(expected_target_revision)):
            raise RuntimeError("コピー先が更新されています。再度コピー操作をしてください")

        source_casts = {}
        for cast in source["voices"]:
            source_casts.setdefault(cast.get("name", ""), []).append(cast)
        for name in cast_names:
            if name not in source_casts or len(source_casts[name]) != 1:
                raise ValueError(f"コピーするキャラクターが存在しないか同名で複数登録されています: {name}")

        used_names = {c.get("name", "") for c in target["voices"]}
        plan, skipped, renamed = [], [], []
        for name in cast_names:
            output_name = name
            if output_name in used_names:
                if conflict == "skip":
                    skipped.append(name)
                    continue
                suffix = 1
                while True:
                    output_name = f"{name}（コピー{suffix if suffix > 1 else ''}）"
                    if output_name not in used_names:
                        break
                    suffix += 1
                renamed.append({"from": name, "to": output_name})
            used_names.add(output_name)
            plan.append((source_casts[name][0], output_name))

        asset_lookup = {a["id"]: a for a in source["assets"]}
        refs_required = []
        for cast, _ in plan:
            for asset_id in cast.get("refs", []):
                if asset_id not in asset_lookup:
                    raise ValueError(f"参照音声が見つかりません: {cast.get('name', '')}")
                if asset_id not in refs_required:
                    refs_required.append(asset_id)

        sources = {}
        for asset_id in refs_required:
            asset = asset_lookup[asset_id]
            filename = str(asset.get("filename", ""))
            if not filename or Path(filename).name != filename or Path(filename).suffix.lower() not in SUPPORTED_AUDIO_EXTS:
                raise ValueError(f"参照音声のファイル名が不正です: {asset_id}")
            source_file = asset_path(source_project_id, asset).resolve()
            if not source_file.is_file():
                raise FileNotFoundError(f"参照音声の実体が見つかりません: {asset.get('name', asset_id)}")
            sources[asset_id] = source_file

        # Create copies before committing the project state; clean up on errors.
        copied_paths = []
        ref_map = {}
        try:
            target_dir = workdir(target_project_id) / asset_subdir("reference")
            target_dir.mkdir(parents=True, exist_ok=True)
            for source_aid, source_file in sources.items():
                source_asset = asset_lookup[source_aid]
                new_aid = uid("asset")
                dest = target_dir / (new_aid + source_file.suffix.lower())
                shutil.copy2(source_file, dest)
                copied_paths.append(dest)
                new_asset = copy.deepcopy(source_asset)
                new_asset.update({"id": new_aid, "filename": dest.name, "kind": "reference"})
                target["assets"].append(new_asset)
                ref_map[source_aid] = new_aid

            for cast, name in plan:
                copied_cast = copy.deepcopy(cast)
                copied_cast["name"] = name
                copied_cast["refs"] = [ref_map[asset_id] for asset_id in cast.get("refs", [])]
                if "id" in copied_cast:
                    copied_cast["id"] = uid("cast")
                target["voices"].append(copied_cast)

            if plan:
                target = save_project(target_project_id, target,
                                      expected_revision=target["revision"], allow_new_assets=True)
        except Exception:
            for path in copied_paths:
                path.unlink(missing_ok=True)
            raise

        return {"project": target, "summary": {
            "copied": len(plan), "skipped": skipped, "renamed": renamed,
            "referencesCopied": len(ref_map),
        }}


def new_line(scene_id, speaker="", text="", caption="", after=0.65):
    return {"id":uid("line"), "sceneId":scene_id, "speaker":speaker,
            "text":text, "caption":caption,"after":after,
            "status":"missing", "takes":[], "selectedTakeId":None, "outputFilename":"",
            "engine":"irodori", "model":"Aratako/Irodori-TTS-v4.1-Small", "mode":"clone", "language":"Japanese", "ttsSource":"cast"}


def attach_take(s, line, asset, generator=None, *, reflow=True):
    if asset.get("kind") == "reference":
        raise ValueError("キャラクター参照音声を台詞へ直接割り当てることはできません")
    take={"id":uid("take"),"assetId":asset["id"],"createdAt":now(),"engine":generator or "import",
          "signature":generation_signature(line, s)}
    line.setdefault("takes",[]).append(take)
    line["selectedTakeId"]=take["id"]
    line["status"]="ready"
    cl=next((c for c in s["clips"] if c.get("lineId")==line["id"] and c["trackId"]=="dialogue"),None)
    if cl is None:
        cl={"id":uid("clip"),"sceneId":line["sceneId"],"trackId":"dialogue","lineId":line["id"],
            "start":0.0, "gain":1.0,"loop":False,"fadeIn":0.0,"fadeOut":0.0,"trimStart":0.0,"pinned":False}
        s["clips"].append(cl)
    cl.update({"assetId":asset["id"],"duration":asset["duration"],"trimStart":0.0})
    if reflow:
        reflow_dialogue(s,line["sceneId"])
    return take


def reflow_dialogue(s,scene_id):
    cursor=0.0
    for line in (x for x in s["lines"] if x["sceneId"]==scene_id):
        clip=next((c for c in s["clips"] if c.get("lineId")==line["id"] and c.get("trackId")=="dialogue"),None)
        if clip:
            if not clip.get("pinned"):
                old_start=float(clip.get("start",0))
                new_start=max(0,cursor)
                clip["start"]=round(new_start,4)
                delta=new_start-old_start
                for related in s["clips"]:
                    if related.get("anchorLineId")==line["id"] and related is not clip:
                        related["start"]=round(max(0,new_start+float(related.get("anchorOffset",0))),4)
            cursor=max(cursor,float(clip["start"])+float(clip.get("duration",0)))
        cursor+=max(0,float(line.get("after",0.65)))


DEFAULT_IRODORI_MODEL = "Aratako/Irodori-TTS-v4.1-Small"


def line_tts_source(line, voice=None):
    """Resolve legacy lines safely: default Small follows the cast, custom models stay explicit."""
    source = line.get("ttsSource")
    if source in ("cast", "line"):
        return "cast" if source == "cast" and voice else "line"
    if not voice:
        return "line"
    # Old script imports created Small settings on *every* line automatically.
    # Do not let those implicit defaults override the speaker's Anime model.
    # Preserve legacy manually-changed nondefault models and engine/mode choices.
    if (line.get("engine", "irodori") != "irodori"
            or line.get("model", DEFAULT_IRODORI_MODEL) not in ("", DEFAULT_IRODORI_MODEL)
            or line.get("mode", "clone") != "clone"):
        return "line"
    return "cast"


def resolve_line_tts(line, voice=None):
    """Return the actual TTS settings used by both preview and generation."""
    voice = voice or {}
    source = line_tts_source(line, voice)
    cfg = voice if source == "cast" else line
    engine = cfg.get("engine") or voice.get("engine") or "irodori"
    if engine not in ("irodori", "qwen"):
        raise ValueError("未対応エンジン: " + str(engine))
    mode = cfg.get("mode") or voice.get("mode") or "clone"
    model = cfg.get("model") or voice.get("model") or (
        DEFAULT_IRODORI_MODEL if engine == "irodori" else "Qwen/Qwen3-TTS-12Hz-1.7B-Base")
    if engine == "qwen" and mode == "design":
        model = "Qwen/Qwen3-TTS-12Hz-1.7B-VoiceDesign"
    language = cfg.get("language") or voice.get("language") or "Japanese"
    return {"source": source, "engine": engine, "mode": mode, "model": model, "language": language}


def generation_signature(line,s):
    import hashlib
    voice=next((v for v in s["voices"] if v["name"]==line.get("speaker")),{})
    data={"text":line.get("text"),"caption":line.get("caption"),
          "effectiveTTS":resolve_line_tts(line,voice)}
    data["profile"]={k:voice.get(k) for k in ("refs","description","refText","language","model","engine","mode")}
    return hashlib.sha256(json.dumps(data,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:24]


def refresh_staleness(s):
    for line in s["lines"]:
        take=next((t for t in line.get("takes",[]) if t.get("id")==line.get("selectedTakeId")),None)
        if not take: line["status"]="missing"
        elif take.get("signature") and take["signature"]!=generation_signature(line,s):
            line["status"]="stale"
        else:line["status"]="ready"


EXCEL_COLUMNS=("台詞ID","シーン","話者","台詞","キャプション","間(秒)","音声ファイル名")

HEADERS={
    "id":{"台詞id","セリフid","voice_id","line_id","id","音声id"},
    "scene":{"シーン","scene","scene_name","場面"},
    "speaker":{"話者","キャラクター","キャラクター名","キャラ","キャラ名","登場人物","役名","speaker","name","character","character_name","人物","発話者"},
    "text":{"台詞","セリフ","本文","text","dialogue","発話"},
    "caption":{"キャプション","演技指示","指示","caption","emotion","ト書き","演出"},
    "after":{"間","間(秒)","間（秒）","次までの間","pause","gap","gap_after","after"},
    "filename":{"音声ファイル名","出力ファイル名","ファイル名","output_filename","outputfilename","wav_filename"},
}

def norm(s):
    return re.sub(r"\s+","",str(s or "").strip()).lower()


def delete_scene(project_id, scene_id, expected_revision):
    """Delete a single scene and its placements, leaving reusable audio untouched.

    A last-scene deletion creates a fresh blank scene for the editor. The old
    project JSON is backed up before saving; asset files and cast profiles are
    not moved or removed.
    """
    with LOCK:
        state = load_project(project_id)
        original_revision = state["revision"]
        if expected_revision is None or int(expected_revision) != original_revision:
            raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
        if not any(scene["id"] == scene_id for scene in state["scenes"]):
            raise FileNotFoundError("削除するシーンが見つかりません")
        old_lines = {line["id"] for line in state["lines"] if line["sceneId"] == scene_id}
        old_clips_count = len(state["clips"])
        state["scenes"] = [scene for scene in state["scenes"] if scene["id"] != scene_id]
        state["lines"] = [line for line in state["lines"] if line["sceneId"] != scene_id]
        state["clips"] = [clip for clip in state["clips"]
                          if clip.get("sceneId") != scene_id
                          and clip.get("lineId") not in old_lines
                          and clip.get("anchorLineId") not in old_lines]
        if not state["scenes"]:
            state["scenes"] = [{"id": uid("scene"), "name": "シーン1"}]
        validate_state(state, load_project(project_id))
        backups = workdir(project_id) / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        backup_name = ("before_scene_delete_" + datetime.now().strftime("%Y%m%d_%H%M%S")
                       + "_" + uuid.uuid4().hex[:8] + ".json")
        shutil.copy2(state_path(project_id), backups / backup_name)
        saved = save_project(project_id, state, expected_revision=original_revision)
        return {"project": saved, "summary": {"removedLines": len(old_lines),
                "removedClips": old_clips_count - len(state["clips"]),
                "backup": backup_name}}


def delete_take_to_trash(project_id, line_id, take_id, expected_revision):
    """Trash one take's sound file only when it has no other consumers.

    Reuse the existing recoverable asset-trash semantics: any fallback take is
    selected and the scene is reflowed. Shared audio is not silently removed.
    """
    with LOCK:
        state = load_project(project_id)
        if expected_revision is None or int(expected_revision) != state["revision"]:
            raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
        line = next((item for item in state["lines"] if item["id"] == line_id), None)
        if line is None:
            raise FileNotFoundError("台詞が見つかりません")
        take = next((item for item in line.get("takes", []) if item["id"] == take_id), None)
        if take is None:
            raise FileNotFoundError("指定したテイクが見つかりません")
        aid = take["assetId"]
        uses = asset_usage(state, aid)
        # Another take (even on the same line), another line, or a cast may
        # depend on the very same file. Never detach them implicitly.
        appearances = sum(t.get("assetId") == aid for l in state["lines"]
                          for t in l.get("takes", []))
        if appearances != 1 or uses["casts"]:
            raise RuntimeError("この音声は他のテイクやキャストでも使用中です。音声ライブラリから管理してください。")
        if any(c.get("lineId") != line_id or c.get("trackId") != "dialogue"
               for c in state["clips"] if c.get("assetId") == aid):
            raise RuntimeError("この音声は別のタイムライン配置でも使用中です。音声ライブラリから管理してください。")
        result = trash_asset(project_id, aid, detach=True,
                             expected_revision=expected_revision)
        result["removedTakeId"] = take_id
        return result


def clear_script(project_id, expected_revision):
    """Remove the imported screenplay/timeline, retaining all media and cast profiles.

    The previous project metadata is backed up before the atomic save. Clearing a
    screenplay never removes physical assets or the reusable character profiles.
    """
    with LOCK:
        state = load_project(project_id)
        revision = state["revision"]
        if int(expected_revision) != revision:
            raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
        removed_lines = len(state["lines"])
        removed_scenes = len(state["scenes"])
        removed_clips = len(state["clips"])
        # Reset to one empty editable scene, not the previously imported scene list.
        state["scenes"] = [{"id": uid("scene"), "name": "シーン1"}]
        state["lines"] = []
        state["clips"] = []
        # Validate the new state before creating a backup or writing to disk.
        validate_state(state, load_project(project_id))
        backups = workdir(project_id) / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        backup_name = ("before_script_clear_" + datetime.now().strftime("%Y%m%d_%H%M%S")
                       + "_" + uuid.uuid4().hex[:8] + ".json")
        shutil.copy2(state_path(project_id), backups / backup_name)
        saved = save_project(project_id, state, expected_revision=revision)
        return {"project": saved, "summary": {"removedLines": removed_lines,
                "removedScenes": removed_scenes, "removedClips": removed_clips,
                "backup": backup_name}}


def import_excel(project_id, stream, mode="merge", expected_revision=None):
    """Import a workbook. 'merge' retains unmatched rows; 'replace' replaces the script.

    Audio assets and character profiles survive either operation. Matching stable
    line IDs keep their takes; records not in a replacement workbook are detached.
    A full JSON snapshot is saved before a replacement commit.
    """
    if mode not in ("merge", "replace"):
        raise ValueError("取り込み方式が正しくありません")
    wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    try:
        with LOCK:
            s = load_project(project_id)
            original_revision = s["revision"]
            if expected_revision is not None and int(expected_revision) != original_revision:
                raise RuntimeError("編集内容の競合を検出しました。再読み込みしてください。")
            old_line_ids = {line["id"] for line in s["lines"]}
            old_scene_ids = {scene["id"] for scene in s["scenes"]}
            old_clips = list(s["clips"])
            sc_map = {scene["name"]: scene["id"] for scene in s["scenes"]}
            old_by_id = {line["id"]: line for line in s["lines"]}
            seen = set()
            warnings = []
            added = updated = matched = 0
            imported = []
            imported_scene_ids = set()
            for ws in wb.worksheets:
                keys = None
                header_row = None
                for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
                    fields = {k: i for i, v in enumerate(row) for k, labels in HEADERS.items() if norm(v) in labels}
                    if "text" in fields:
                        header_row, keys = row_num, fields
                        break
                if not keys:
                    if mode == "replace":
                        raise ValueError(f"{ws.title}: 台詞列がありません。全件置き換えを中止しました")
                    warnings.append(f"{ws.title}: 台詞列が見つからずスキップしました")
                    continue

                def cell(row, key):
                    i = keys.get(key)
                    return row[i] if i is not None and i < len(row) else None

                for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                    txt = str(cell(row, "text") or "").strip()
                    if not txt:
                        continue
                    speaker = str(cell(row, "speaker") or "ナレーション").strip()
                    scene_name = str(cell(row, "scene") or ws.title).strip()
                    if scene_name not in sc_map:
                        sid = uid("scene")
                        s["scenes"].append({"id": sid, "name": scene_name})
                        sc_map[scene_name] = sid
                    sid = sc_map[scene_name]
                    imported_scene_ids.add(sid)
                    external_id = str(cell(row, "id") or "").strip()
                    line = old_by_id.get(external_id) if external_id else None
                    if external_id and external_id in seen:
                        if mode == "replace":
                            raise ValueError(f"{ws.title} {row_no}行目: 台詞IDが重複しています ({external_id})")
                        warnings.append(f"{ws.title} {row_no}行目: ID重複のためスキップ ({external_id})")
                        continue
                    if line and line["sceneId"] != sid:
                        # Keep sound effects anchored to a retained dialogue line in the right scene.
                        for clip in s["clips"]:
                            if clip.get("lineId") == line["id"] or clip.get("anchorLineId") == line["id"]:
                                clip["sceneId"] = sid
                    if not line and not external_id:
                        possibilities = [x for x in s["lines"] if x["id"] not in seen
                                         and x["sceneId"] == sid and x["speaker"] == speaker and x["text"] == txt]
                        if len(possibilities) == 1:
                            line = possibilities[0]
                            matched += 1
                    if not line:
                        if external_id and external_id in old_by_id:
                            warnings.append(f"{ws.title} {row_no}行目: IDの照合に失敗")
                        line = new_line(sid)
                        if external_id and re.fullmatch(r"line_[a-f0-9]{16}", external_id):
                            line["id"] = external_id
                        s["lines"].append(line)
                        old_by_id[line["id"]] = line
                        added += 1
                    else:
                        updated += 1
                    line["speaker"] = speaker
                    line["text"] = txt
                    line["sceneId"] = sid
                    if "filename" in keys:
                        line["outputFilename"] = str(cell(row,"filename") or "").strip()
                    if "caption" in keys:
                        line["caption"] = str(cell(row, "caption") or "")
                    if "after" in keys and cell(row, "after") not in (None, ""):
                        try:
                            line["after"] = max(0, float(cell(row, "after")))
                        except (TypeError, ValueError):
                            warnings.append(f"{row_no}行目: 間の数値を読み取れません")
                    seen.add(line["id"])
                    imported.append(line["id"])
            # Remove the empty default scene when the Excel declares its own scenes.
            # Otherwise a cleared project shows a misleading extra "シーン1" tab.
            if (imported and not old_line_ids and not old_clips
                    and len(old_scene_ids) == 1 and len(s["scenes"]) > 1):
                empty_defaults = {x["id"] for x in s["scenes"]
                                  if x["name"] == "シーン1" and x["id"] not in imported_scene_ids}
                s["scenes"] = [x for x in s["scenes"] if x["id"] not in empty_defaults]
            if mode == "replace" and not imported:
                raise ValueError("読み込み可能な台詞が0件のため全件置き換えを中止しました")
            if mode == "replace":
                # 'Workbook is the complete script': remove all lines/scenes not represented.
                s["lines"] = [line for line in s["lines"] if line["id"] in seen]
                s["scenes"] = [scene for scene in s["scenes"] if scene["id"] in imported_scene_ids]
                # Keep independent BGM/SFX clips in retained scenes, but discard
                # orphaned voice clips / anchored effects and old scenes' clips.
                s["clips"] = [clip for clip in s["clips"]
                              if clip.get("sceneId") in imported_scene_ids
                              and (not clip.get("lineId") or clip["lineId"] in seen)
                              and (not clip.get("anchorLineId") or clip["anchorLineId"] in seen)]
                order = {line_id: i for i, line_id in enumerate(imported)}
                s["lines"].sort(key=lambda line: order[line["id"]])
                s["scenes"].sort(key=lambda scene: min(
                    (order[line["id"]] for line in s["lines"] if line["sceneId"] == scene["id"]),
                    default=len(imported)))
            else:
                # Non-destructive merge: keep untouched entries after imported rows.
                order = {line_id: i for i, line_id in enumerate(imported)}
                s["lines"].sort(key=lambda line: (0, order[line["id"]]) if line["id"] in order else (1, 0))
            known = {voice.get("name") for voice in s["voices"]}
            for line in s["lines"]:
                speaker = line.get("speaker", "")
                if speaker and speaker not in known:
                    s["voices"].append({"name": speaker, "engine": "irodori",
                        "model": "Aratako/Irodori-TTS-v4.1-Small", "mode": "clone",
                        "description": "", "language": "Japanese", "refText": "", "refs": []})
                    known.add(speaker)
            for scene in s["scenes"]:
                reflow_dialogue(s, scene["id"])
            refresh_staleness(s)
            removed = len(old_line_ids - {line["id"] for line in s["lines"]})
            removed_scenes = len(old_scene_ids - {scene["id"] for scene in s["scenes"]})
            removed_clips = len(old_clips) - len(s["clips"])
            backup_name = None
            if mode == "replace":
                # Save a full project metadata snapshot before modifying the state.
                # Audio bytes are never removed by this operation.
                backups = workdir(project_id) / "backups"
                backups.mkdir(parents=True, exist_ok=True)
                backup_name = f"before_excel_replace_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}.json"
                shutil.copy2(state_path(project_id), backups / backup_name)
            s = save_project(project_id, s, expected_revision=original_revision)
            return {"project": s, "summary": {"added": added, "updated": updated,
                    "matched": matched, "removed": removed, "removedScenes": removed_scenes,
                    "removedClips": max(0, removed_clips), "backup": backup_name,
                    "mode": mode, "warnings": warnings}}
    finally:
        wb.close()


def excel_template():
    """An empty, import-ready script template with guidance outside input columns.

    The importer only uses the named A:G columns. Never put sample dialogue
    into A:G: otherwise a first-time user might import it as a real line.
    """
    from io import BytesIO
    from openpyxl.styles import Alignment, Font, PatternFill

    wb=openpyxl.Workbook()
    ws=wb.active
    ws.title="台本"
    ws.append(list(EXCEL_COLUMNS))
    ws.freeze_panes="A2"
    ws.auto_filter.ref="A1:G1"
    for cell in ws[1][:7]:
        cell.font=Font(color="FFFFFF",bold=True)
        cell.fill=PatternFill("solid",fgColor="1C3347")
        cell.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=27
    for col,width in {"A":24,"B":20,"C":18,"D":64,"E":56,"F":14,"G":32,"H":4,"I":26,"J":64}.items():
        ws.column_dimensions[col].width=width
    guides=[
        ("入力ガイド（取り込み対象外）","左側のA〜G列に2行目から入力してください。ここは記入不要です。"),
        ("話者（必須）","例：真歩流 / 愛理。キャラクター名を入力します。"),
        ("台詞（必須）","例：おかえり、お姉ちゃん。"),
        ("シーン（任意）","例：第一幕。空欄ならシート名の「台本」がシーン名になります。"),
        ("キャプション（任意）","例：驚いて、小さな声で話す。"),
        ("間(秒)（任意）","例：0.6。次のセリフまでの間を秒で指定します。"),
        ("台詞ID（初回空欄）","最初は空欄でOK。再取り込みはアプリの「台本をExcelで書き出す」でID付きのファイルを使ってください。"),
        ("音声ファイル名（任意）","例：scene_001.wav。空欄なら自動命名されます。"),
        ("取り込み方法","A〜G列に台本を入力し、.xlsxで保存して「Excel取込」から読み込みます。"),
        ("注意","I・J列の説明は取り込み対象外です。A〜G列にサンプル台詞は入っていません。"),
    ]
    for row,(label,description) in enumerate(guides,1):
        item=ws.cell(row=row,column=9,value=label)
        item.font=Font(bold=True,color="FFFFFF" if row==1 else "1C3347")
        item.fill=PatternFill("solid",fgColor="1C3347" if row==1 else "E6F2F7")
        item.alignment=Alignment(vertical="center",wrap_text=True)
        detail=ws.cell(row=row,column=10,value=description)
        detail.alignment=Alignment(vertical="center",wrap_text=True)
        if row==1:
            detail.font=Font(bold=True,color="FFFFFF")
            detail.fill=PatternFill("solid",fgColor="1C3347")
        if row>1: ws.row_dimensions[row].height=35
    # Text input columns are formatted for easy reading without adding data.
    for row in ws.iter_rows(min_row=2,max_row=40,min_col=1,max_col=7):
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True)
    buf=BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_excel(project_id):
    from io import BytesIO
    s=load_project(project_id)
    wb=openpyxl.Workbook()
    wb.remove(wb.active)
    for scene in s["scenes"]:
        # Excel sheet names must be unique and <=31 chars.
        name=re.sub(r'[\\/*?:\[\]]','_',scene["name"])[:28] or "Scene"
        base=name;i=1
        while name in wb.sheetnames:
            i+=1;name=base[:25]+f"_{i}"
        ws=wb.create_sheet(name)
        ws.append(list(EXCEL_COLUMNS))
        for line in (l for l in s["lines"] if l["sceneId"]==scene["id"]):
            ws.append([line["id"],scene["name"],line["speaker"],line["text"],line.get("caption",""),line.get("after",0.65),line.get("outputFilename","")])
        ws.freeze_panes="A2"
        from openpyxl.styles import Font,PatternFill,Alignment
        for c in ws[1]:c.font=Font(color="FFFFFF",bold=True);c.fill=PatternFill("solid",fgColor="1C3347")
        for col,width in {"A":24,"B":20,"C":18,"D":64,"E":56,"F":14,"G":32}.items():ws.column_dimensions[col].width=width
        for row in ws.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)
    buf=BytesIO();wb.save(buf);buf.seek(0);return buf


def register_asset(project_id,src_path,name,kind="other"):
    with LOCK:
        s=load_project(project_id)
        ext=src_path.suffix.lower()
        if ext not in (".wav",".mp3",".flac",".ogg",".m4a",".aac",".opus"):
            raise ValueError("WAV / MP3 / FLAC / OGG / M4A / AAC / OPUSのみ対応")
        asset_id=uid("asset")
        dest=workdir(project_id)/asset_subdir(kind)/(asset_id+ext)
        dest.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(src_path,dest)
        from render import duration_of
        try:duration=duration_of(dest)
        except Exception:
            dest.unlink(missing_ok=True)
            raise
        meta={"id":asset_id,"filename":dest.name,"name":name[:160],"kind":kind,"duration":round(duration,4)}
        s["assets"].append(meta)
        s=save_project(project_id,s,allow_new_assets=True)
        return {"asset":meta,"project":s}


SUPPORTED_AUDIO_EXTS = (".wav", ".mp3", ".flac", ".ogg", ".m4a", ".aac", ".opus")


def import_audio_batch(project_id, items, kind="voice", line_ids=None, cast_name=""):
    """Import multiple uploaded audio files in one project revision.

    `items`: list of (temporary file path, original display name), in UI order.
    `line_ids`: one existing line ID or None for each item (None means library only).
    Existing takes are kept; new assets get unique IDs. Invalid requests do not
    change project metadata and uploaded/copied files are cleaned up on error.
    """
    if not 1 <= len(items) <= 250:
        raise ValueError("一度に取り込める音声は1〜250件です")
    if kind not in {"voice", "reference", "sfx", "bgm", "ambience", "other"}:
        raise ValueError("音声の種類が不正です")
    if line_ids is None:
        line_ids = [None] * len(items)
    if not isinstance(line_ids, list) or len(line_ids) != len(items):
        raise ValueError("音声と割り当て先の件数が一致しません")
    line_ids = [x or None for x in line_ids]
    if len([x for x in line_ids if x]) != len(set(x for x in line_ids if x)):
        raise ValueError("同じ台詞に複数の音声を一括割り当てできません")
    if any(not Path(name.replace('\\', '/')).suffix.lower() in SUPPORTED_AUDIO_EXTS for _, name in items):
        raise ValueError("WAV / MP3 / FLAC / OGG / M4A / AAC / OPUSのみ対応")
    # Parse duration and verify all files before placing anything in workspace.
    from render import duration_of
    durations = []
    for path, _ in items:
        duration = duration_of(path)
        if not 0 < duration <= 7200:
            raise ValueError("音声ファイルの長さが不正です（2時間以内）")
        durations.append(duration)

    with LOCK:
        state = load_project(project_id)
        lines = {line['id']: line for line in state['lines']}
        for line_id in line_ids:
            if line_id and line_id not in lines:
                raise ValueError(f"割り当て先の台詞が見つかりません: {line_id}")
        voice = None
        if cast_name:
            if kind != 'reference':
                raise ValueError("キャストへの登録には参照音声モードを指定してください")
            voice = next((v for v in state['voices'] if v['name'] == cast_name), None)
            if voice is None:
                raise ValueError(f"キャストが見つかりません: {cast_name}")
        if kind == 'reference' and any(line_ids):
            raise ValueError("参照音声は台詞へ直接割り当てできません")
        new_assets = []
        copied = []
        scenes = set()
        try:
            for (path, name), duration, line_id in zip(items, durations, line_ids):
                ext = Path(name.replace('\\', '/')).suffix.lower()
                aid = uid('asset')
                dest = workdir(project_id) / asset_subdir(kind) / (aid + ext)
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, dest)
                copied.append(dest)
                asset = {"id": aid, "filename": dest.name, "name": Path(name.replace('\\','/')).name[:160],
                         "kind": kind, "duration": round(duration, 4)}
                new_assets.append(asset)
                state['assets'].append(asset)
                if line_id:
                    line = lines[line_id]
                    attach_take(state, line, asset, reflow=False)
                    scenes.add(line['sceneId'])
                if voice is not None:
                    voice.setdefault('refs', []).append(aid)
            for scene_id in scenes:
                reflow_dialogue(state, scene_id)
            state = save_project(project_id, state, allow_new_assets=True)
        except Exception:
            for dest in copied:
                dest.unlink(missing_ok=True)
            raise
    return {"project": state, "assets": new_assets,
            "summary": {"imported": len(new_assets), "assigned": sum(bool(x) for x in line_ids),
                        "references": len(new_assets) if cast_name else 0}}
