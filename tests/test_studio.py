import io
import sys
from pathlib import Path
import numpy as np
import pytest
import soundfile as sf
from fastapi.testclient import TestClient
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import core
from app import app

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(core,'WORKSPACE',tmp_path)
    monkeypatch.setattr(__import__('importlib').import_module('app'), 'EXPORTS_DIR', tmp_path / 'exports')
    return TestClient(app)


def project(client):
    r=client.post('/api/projects',json={'title':'舞黒館テスト'})
    assert r.status_code==200,r.text
    return r.json()


def wav_blob(freq=330,seconds=.3):
    sr=44100;t=np.arange(int(sr*seconds))/sr
    buf=io.BytesIO();sf.write(buf,(np.sin(2*np.pi*freq*t)*.15).astype('float32'),sr,format='WAV');buf.seek(0)
    return buf


def upload(client,pid,freq=440,kind='voice'):
    r=client.post(f'/api/projects/{pid}/assets',files={'file':('voice.wav',wav_blob(freq),'audio/wav')},data={'kind':kind})
    assert r.status_code==200,r.text
    return r.json()['asset']


def test_project_audio_workflow_and_mixer(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    l=core.new_line(sid,'真歩流','サブレがない！','驚き',.4)
    s['lines']=[l]
    r=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    a=upload(client,pid,440,'voice')
    r=client.post(f'/api/projects/{pid}/lines/{l["id"]}/attach/{a["id"]}')
    assert r.status_code==200,r.text
    s=r.json();assert s['lines'][0]['selectedTakeId'] and s['lines'][0]['status']=='ready'
    s['clips'].append({'id':core.uid('clip'),'sceneId':sid,'trackId':'bgm','assetId':a['id'],
          'start':0,'duration':.7,'loop':True,'gain':.25,'fadeIn':.1,'fadeOut':.1})
    r=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    r=client.post(f'/api/projects/{pid}/render',json={'sceneId':sid,'format':'wav'})
    assert r.status_code==200,r.text
    path=r.json()['url'];download=client.get(path);assert download.status_code==200
    audio,samplerate=sf.read(io.BytesIO(download.content));assert samplerate==44100
    assert len(audio)>=int(.69*samplerate) and np.abs(audio[:int(.28*samplerate)]).mean()>.01
    assert client.get(f'/api/projects/{pid}').json()['clips'][1]['loop']


def test_excel_roundtrip_preserves_ids_and_stale_audio(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    l=core.new_line(sid,'愛理','どこに行ったの？','少し困惑',.5)
    s['lines']=[l]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    a=upload(client,pid,200)
    s=client.post(f'/api/projects/{pid}/lines/{l["id"]}/attach/{a["id"]}').json()
    x=client.get(f'/api/projects/{pid}/excel/export');assert x.status_code==200
    wb=openpyxl.load_workbook(io.BytesIO(x.content));ws=wb.active
    assert ws['A2'].value==l['id']
    ws['D2']='ねえ、サブレはどこに行ったの？'
    ws.append([None,'シーン1','真歩流','分からないわ','落ち着いて',.7])
    bio=io.BytesIO();wb.save(bio);bio.seek(0)
    r=client.post(f'/api/projects/{pid}/excel/import',files={'file':('script.xlsx',bio,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert r.status_code==200,r.text
    s=r.json()['project'];assert len(s['lines'])==2
    prior=next(x for x in s['lines'] if x['id']==l['id']);assert prior['status']=='stale'
    assert prior['takes'][0]['assetId']==a['id']
    assert r.json()['summary']['added']==1
    # Reimport the same ID-bearing Excel must not duplicate known lines.
    x=client.get(f'/api/projects/{pid}/excel/export')
    r=client.post(f'/api/projects/{pid}/excel/import',files={'file':('again.xlsx',x.content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert len(r.json()['project']['lines'])==2


def test_audio_asset_security_and_revision_conflict(client):
    s=project(client);pid=s['id']
    r=client.post(f'/api/projects/{pid}/assets',files={'file':('notes.txt',b'audio','text/plain')},data={'kind':'sfx'})
    assert r.status_code==400
    malicious=dict(s);malicious['assets']=[{'id':'asset_bad','filename':'../../x.wav','duration':1}]
    r=client.put(f'/api/projects/{pid}',json={'state':malicious,'expectedRevision':s['revision']})
    assert r.status_code==400
    s['title']='Updated'
    r=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert r.status_code==200
    r=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert r.status_code==409


def test_timeline_anchor_reflow():
    s=core.default_state('test','p_1111111111111111');sid=s['scenes'][0]['id']
    l=core.new_line(sid,'a','hello',after=1)
    a={'id':'asset_1','filename':'a.wav','duration':2};s['assets'].append(a);s['lines'].append(l)
    core.attach_take(s,l,a)
    c=s['clips'][0]
    s['clips'].append({'id':'clip_sfx','sceneId':sid,'trackId':'sfx','assetId':'asset_1','start':.3,'duration':1,
          'anchorLineId':l['id'],'anchorOffset':.3})
    # Move the dialogue start manually, reflow does not override pinned clips.
    c['pinned']=True;c['start']=3
    core.reflow_dialogue(s,sid)
    assert c['start']==3
    c['pinned']=False
    core.reflow_dialogue(s,sid)
    assert c['start']==0 and s['clips'][1]['start']==.3


def test_generate_job_registration_with_fake_tts(client,monkeypatch,tmp_path):
    import json
    import app as api
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['lines'].append(core.new_line(sid,'真歩流','ここはどこ？','困惑',.4))
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    # Mock only the external model runtime, not project or audio handling.
    monkeypatch.setattr(api,'python_for',lambda *_:sys.executable)
    def fake_subprocess_run(cmd,**kwargs):
        job=json.loads(Path(cmd[-2]).read_text(encoding='utf-8'))
        rows=[]
        for item in job['items']:
            path=Path(item['output']);path.parent.mkdir(exist_ok=True)
            sf.write(path,np.zeros(4410,dtype='float32'),44100)
            rows.append({'id':item['id'],'ok':True,'output':str(path)})
        Path(cmd[-1]).write_text(json.dumps({'ok':True,'items':rows}),encoding='utf-8')
        return type('Finished',(),{'returncode':0})()
    monkeypatch.setattr(api.subprocess,'run',fake_subprocess_run)
    job_id=core.uid('job')
    api.JOBS[job_id]={'status':'running','completed':0,'total':0,'errors':[],'message':''}
    api.run_generation(job_id,pid,[s['lines'][0]['id']],40)
    assert api.JOBS[job_id]['status']=='done'
    final=client.get(f'/api/projects/{pid}').json()
    assert len(final['lines'][0]['takes'])==1 and len(final['clips'])==1
    assert final['lines'][0]['status']=='ready'


def test_voice_design_creates_reusable_reference(client,monkeypatch):
    import json
    import app as api
    s=project(client);pid=s['id']
    s['voices']=[{'name':'真歩流','engine':'qwen','model':'Qwen/Qwen3-TTS-12Hz-1.7B-VoiceDesign',
                 'mode':'design','description':'明るく元気な少女の声','language':'Japanese','refText':'','refs':[]}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    monkeypatch.setattr(api,'python_for',lambda *_:sys.executable)
    def fake_run(cmd,**kwargs):
        job=json.loads(Path(cmd[-2]).read_text(encoding='utf-8'))
        item=job['items'][0];assert job['model'].endswith('VoiceDesign')
        assert item['mode']=='design'
        f=Path(item['output']);sf.write(f,np.zeros(4410,dtype='float32'),44100)
        Path(cmd[-1]).write_text(json.dumps({'ok':True,'items':[{'id':item['id'],'ok':True,'output':str(f)}]}))
        return type('Completed',(),{'returncode':0})()
    monkeypatch.setattr(api.subprocess,'run',fake_run)
    jid=core.uid('job');api.JOBS[jid]={'status':'running','total':1,'completed':0,'errors':[],'message':''}
    api.run_voice_design(jid,pid,'真歩流','こんにちは！')
    assert api.JOBS[jid]['status']=='done'
    latest=client.get(f'/api/projects/{pid}').json()
    voice=latest['voices'][0]
    assert len(voice['refs'])==1 and voice['refText']=='こんにちは！'
    assert voice['mode']=='clone' and voice['model'].endswith('Base')
    assert latest['assets'][0]['kind']=='reference'


def test_batch_upload_assigns_in_file_order_and_keeps_old_takes(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    lines=[core.new_line(sid,'真歩流',f'台詞{i}',after=.2) for i in range(1,4)]
    s['lines']=lines
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    old=upload(client,pid,freq=180)
    s=client.post(f'/api/projects/{pid}/lines/{lines[0]["id"]}/attach/{old["id"]}').json()
    old_take_id=s['lines'][0]['selectedTakeId']
    r=client.post(f'/api/projects/{pid}/assets/batch',
                  files=[('files',('voice_01.wav',wav_blob(300),'audio/wav')),
                         ('files',('voice_02.wav',wav_blob(420),'audio/wav')),
                         ('files',('voice_10.wav',wav_blob(540),'audio/wav'))],
                  data={'kind':'voice','lineIds':__import__('json').dumps([l['id'] for l in lines])})
    assert r.status_code==200,r.text
    body=r.json();s=body['project']
    assert body['summary']=={'imported':3,'assigned':3,'references':0}
    assert [a['name'] for a in body['assets']]==['voice_01.wav','voice_02.wav','voice_10.wav']
    assert [next(t for t in l['takes'] if t['id']==l['selectedTakeId'])['assetId'] for l in s['lines']]==[a['id'] for a in body['assets']]
    assert len(s['lines'][0]['takes'])==2 and s['lines'][0]['takes'][0]['id']==old_take_id
    assert [c['lineId'] for c in s['clips']]==[l['id'] for l in lines]
    assert [round(c['start'],2) for c in s['clips']]==[0,.5,1.0]
    assert len(list((core.workdir(pid)/'dialogue_audio').glob('*.wav')))==4
    assert s['revision']==5 # create, save lines, upload original, attach original, entire batch
    rr=client.post(f'/api/projects/{pid}/render',json={'sceneId':sid,'format':'wav'})
    assert rr.status_code==200,rr.text


def test_batch_upload_skips_targets_and_can_register_cast_references(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['voices']=[{'name':'愛理','refs':[], 'engine':'irodori', 'mode':'clone', 'model':'test'}]
    s['lines']=[core.new_line(sid,'愛理','よし')]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    import json
    r=client.post(f'/api/projects/{pid}/assets/batch',
                 files=[('files',('a.wav',wav_blob(),'audio/wav')),('files',('b.wav',wav_blob(),'audio/wav'))],
                 data={'kind':'reference','castName':'愛理','lineIds':json.dumps([None,None])})
    assert r.status_code==200,r.text
    res=r.json();assert res['summary']=={'imported':2,'assigned':0,'references':2}
    assert res['project']['voices'][0]['refs']==[a['id'] for a in res['assets']]
    assert res['project']['lines'][0]['selectedTakeId'] is None
    assert not res['project']['clips']
    r=client.post(f'/api/projects/{pid}/assets/batch',
                 files=[('files',('free.wav',wav_blob(),'audio/wav'))],
                 data={'kind':'voice','lineIds':json.dumps([None])})
    assert r.status_code==200,r.text
    assert r.json()['summary']['assigned']==0


def test_batch_upload_rejects_duplicate_targets_and_invalid_files_atomically(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    l=core.new_line(sid,'真歩流','こんにちは')
    s['lines'].append(l)
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    import json
    for files,targets,kind,cast in [
        ([('files',('a.wav',wav_blob(),'audio/wav')),('files',('b.wav',wav_blob(),'audio/wav'))],[l['id'],l['id']],'voice',''),
        ([('files',('a.wav',wav_blob(),'audio/wav')),('files',('fake.txt',b'junk','text/plain'))],[l['id'],None],'voice',''),
        ([('files',('a.wav',wav_blob(),'audio/wav'))],[l['id']],'reference',''),
        ([('files',('a.wav',wav_blob(),'audio/wav'))],[None],'reference','誰もいない'),
        ([('files',('broken.wav',b'fake-wav','audio/wav'))],[None],'voice','')
    ]:
        r=client.post(f'/api/projects/{pid}/assets/batch',files=files,data={
            'kind':kind,'lineIds':json.dumps(targets),'castName':cast})
        assert r.status_code==400,r.text
        assert client.get(f'/api/projects/{pid}').json()['revision']==s['revision']
        assert not list((core.workdir(pid)/'assets').iterdir())


def test_create_multiple_projects_and_ui_create_handler(client):
    first=project(client)
    second=client.post('/api/projects',json={'title':'新しいドラマ'}).json()
    assert first['id']!=second['id']
    assert {p['id'] for p in client.get('/api/projects').json()}=={first['id'],second['id']}
    assert second['lines']==[] and second['assets']==[]
    js=Path(__file__).resolve().parents[1]/'static'/'app.js'
    content=js.read_text(encoding='utf-8')
    assert "target.id==='createProjectBtn'" in content
    assert 'createProject(title).then(' in content


def test_cast_references_are_separate_from_dialogue_assets(client):
    state=project(client);pid=state['id'];sid=state['scenes'][0]['id']
    dialogue=upload(client,pid,kind='voice')
    reference=upload(client,pid,kind='reference')
    state=client.get(f'/api/projects/{pid}').json()
    state['voices']=[{'name':'真歩流','refs':[dialogue['id']]}]
    rejected=client.put(f'/api/projects/{pid}',json={'state':state,'expectedRevision':state['revision']})
    assert rejected.status_code==400
    assert '参照音声' in rejected.text
    state['voices'][0]['refs']=[reference['id']]
    assert client.put(f'/api/projects/{pid}',json={'state':state,'expectedRevision':state['revision']}).status_code==200
    # Reference audio cannot be attached to spoken lines.
    state=client.get(f'/api/projects/{pid}').json()
    line=core.new_line(sid,'真歩流','こんにちは')
    state['lines'].append(line)
    state=client.put(f'/api/projects/{pid}',json={'state':state,'expectedRevision':state['revision']}).json()
    r=client.post(f'/api/projects/{pid}/lines/{line["id"]}/attach/{reference["id"]}')
    assert r.status_code==400
    assert len(client.get(f'/api/projects/{pid}').json()['lines'][0]['takes'])==0


def test_unused_asset_trash_and_restore(client):
    s=project(client);pid=s['id'];asset=upload(client,pid)
    original=core.asset_path(pid, asset)
    assert original.is_file()
    response=client.delete(f'/api/projects/{pid}/assets/{asset["id"]}')
    assert response.status_code==200,response.text
    assert not original.exists()
    assert (core.workdir(pid)/'trash'/asset['filename']).is_file()
    assert asset['id'] not in [x['id'] for x in response.json()['project']['assets']]
    assert client.get(f'/api/projects/{pid}/trash/assets').json()[0]['asset']['id']==asset['id']
    restored=client.post(f'/api/projects/{pid}/trash/assets/{asset["id"]}/restore')
    assert restored.status_code==200,restored.text
    assert original.is_file()
    assert asset['id'] in [x['id'] for x in restored.json()['project']['assets']]
    assert client.get(f'/api/projects/{pid}/trash/assets').json()==[]


def test_linked_asset_requires_explicit_detach_and_preserves_other_takes(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['lines']=[core.new_line(sid,'愛理','確かにあるよ')]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    one=upload(client,pid,freq=222);two=upload(client,pid,freq=444)
    s=client.post(f'/api/projects/{pid}/lines/{s["lines"][0]["id"]}/attach/{one["id"]}').json()
    s=client.post(f'/api/projects/{pid}/lines/{s["lines"][0]["id"]}/attach/{two["id"]}').json()
    usage=client.get(f'/api/projects/{pid}/assets/{two["id"]}/usage').json()
    assert len(usage['lines'])==1 and len(usage['clips'])==1
    assert client.delete(f'/api/projects/{pid}/assets/{two["id"]}').status_code==409
    after=client.get(f'/api/projects/{pid}').json()
    assert len(after['lines'][0]['takes'])==2
    deleted=client.delete(f'/api/projects/{pid}/assets/{two["id"]}?detach=true')
    assert deleted.status_code==200,deleted.text
    after=deleted.json()['project'];line=after['lines'][0]
    assert len(line['takes'])==1 and line['takes'][0]['assetId']==one['id']
    assert line['selectedTakeId']==line['takes'][0]['id']
    assert after['clips'][0]['assetId']==one['id']
    assert (core.workdir(pid)/'trash'/two['filename']).exists()
    restored=client.post(f'/api/projects/{pid}/trash/assets/{two["id"]}/restore').json()['project']
    assert len(restored['lines'][0]['takes'])==1 # Re-association is intentionally manual.


def test_reference_removal_is_guarded_and_invalidates_voice_profile(client):
    s=project(client);pid=s['id'];reference=upload(client,pid,kind='reference')
    s=client.get(f'/api/projects/{pid}').json()
    s['voices']=[{'name':'真歩流','refs':[reference['id']],'engine':'irodori'}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    usage=client.get(f'/api/projects/{pid}/assets/{reference["id"]}/usage').json()
    assert usage['casts']==['真歩流']
    assert client.delete(f'/api/projects/{pid}/assets/{reference["id"]}').status_code==409
    result=client.delete(f'/api/projects/{pid}/assets/{reference["id"]}?detach=true')
    assert result.status_code==200,result.text
    assert result.json()['project']['voices'][0]['refs']==[]


def test_legacy_voice_reference_remains_removable(client):
    # An old project can contain dialogue-kind audio under refs; do not break
    # loading/editing until the user removes the old checkbox.
    s=project(client);pid=s['id'];legacy=upload(client,pid,kind='voice')
    s=client.get(f'/api/projects/{pid}').json()
    s['voices']=[{'name':'真歩流','refs':[legacy['id']]}]
    # Persist a legacy state as though created by pre-v0.1.3 app.
    (core.workdir(pid)/'project.json').write_text(__import__('json').dumps(s,ensure_ascii=False),encoding='utf-8')
    s['title']='保存しても既存の選択肢は維持される'
    res=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert res.status_code==200,res.text
    s=res.json();s['voices'][0]['refs']=[]
    assert client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).status_code==200


def test_bulk_trash_unused_only_and_rejects_linked(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    a=upload(client,pid,freq=222);b=upload(client,pid,freq=333)
    response=client.post(f'/api/projects/{pid}/assets/bulk/trash',json={'assetIds':[a['id'],b['id']]})
    assert response.status_code==200,response.text
    assert len(response.json()['trashed'])==2
    assert client.get(f'/api/projects/{pid}').json()['assets']==[]
    assert len(client.get(f'/api/projects/{pid}/trash/assets').json())==2
    another=upload(client,pid,freq=444)
    s=client.get(f'/api/projects/{pid}').json()
    line=core.new_line(sid,'真歩流','聞こえる？')
    s['lines'].append(line)
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    s=client.post(f'/api/projects/{pid}/lines/{line["id"]}/attach/{another["id"]}').json()
    extra=upload(client,pid,freq=555)
    response=client.post(f'/api/projects/{pid}/assets/bulk/trash',json={'assetIds':[another['id'],extra['id']]})
    assert response.status_code==409
    current=client.get(f'/api/projects/{pid}').json()
    assert {x['id'] for x in current['assets']}=={another['id'],extra['id']}


def test_copy_selected_casts_with_independent_reference_assets(client):
    """Multi-select copies settings/files, preserves sharing, and source remains intact."""
    source = project(client)
    target = client.post('/api/projects', json={'title': '続編'}).json()
    pid, tid = source['id'], target['id']
    shared = upload(client, pid, 180, 'reference')
    separate = upload(client, pid, 440, 'reference')
    source = client.get(f'/api/projects/{pid}').json()
    source['voices'] = [
        {'name': '真歩流', 'engine': 'irodori', 'model': 'Aratako/Irodori-TTS-v4.1-Anime',
         'mode': 'clone', 'language': 'Japanese', 'description': '若い元気な探偵',
         'refText': '', 'refs': [shared['id'], separate['id']]},
        {'name': '愛理', 'engine': 'qwen', 'model': 'Qwen/Qwen3-TTS-12Hz-1.7B-Base',
         'mode': 'clone', 'language': 'Japanese', 'description': '落ち着いた妹',
         'refText': '私は愛理です。', 'refs': [shared['id']]},
        {'name': '零度', 'engine': 'irodori', 'model': 'test', 'mode': 'design',
         'description': '渋い', 'refs': []},
    ]
    source = client.put(f'/api/projects/{pid}', json={
        'state': source, 'expectedRevision': source['revision']
    }).json()
    resp = client.post(f'/api/projects/{pid}/casts/copy', json={
        'targetProjectId': tid, 'castNames': ['真歩流', '愛理'],
        'expectedSourceRevision': source['revision'],
        'expectedTargetRevision': target['revision'],
    })
    assert resp.status_code == 200, resp.text
    payload = resp.json()
    dest = payload['project']
    assert payload['summary'] == {'copied': 2, 'skipped': [], 'renamed': [], 'referencesCopied': 2}
    assert [v['name'] for v in dest['voices']] == ['真歩流', '愛理']
    assert dest['voices'][0]['description'] == '若い元気な探偵'
    assert dest['voices'][1]['refText'] == '私は愛理です。'
    assert dest['lines'] == [] and dest['clips'] == []
    assert len(dest['assets']) == 2
    copied_shared = dest['voices'][0]['refs'][0]
    assert dest['voices'][1]['refs'] == [copied_shared]
    assert copied_shared != shared['id']
    assert all(a['kind'] == 'reference' for a in dest['assets'])
    by_id = {a['id']: a for a in dest['assets']}
    dest_file = core.asset_path(tid, by_id[copied_shared])
    src_file = core.asset_path(pid, shared)
    assert dest_file.read_bytes() == src_file.read_bytes()
    assert client.delete(f'/api/projects/{pid}/assets/{shared["id"]}?detach=true').status_code == 200
    assert dest_file.exists()
    assert client.get(f'/api/projects/{tid}').json()['voices'][1]['refs'] == [copied_shared]


def test_copy_single_cast_collision_skip_and_rename(client):
    src = project(client)
    tgt = client.post('/api/projects', json={'title': '別作品'}).json()
    src['voices'] = [{'name': '真歩流', 'refs': [], 'description': '探偵'},
                     {'name': '愛理', 'refs': [], 'description': '妹'}]
    tgt['voices'] = [{'name': '真歩流', 'refs': [], 'description': '変更禁止'},
                     {'name': '愛理', 'refs': [], 'description': '既存の愛理'},
                     {'name': '愛理（コピー）', 'refs': [], 'description': '既存'}]
    src = client.put(f'/api/projects/{src["id"]}', json={'state': src, 'expectedRevision': src['revision']}).json()
    tgt = client.put(f'/api/projects/{tgt["id"]}', json={'state': tgt, 'expectedRevision': tgt['revision']}).json()
    req = {'targetProjectId': tgt['id'], 'castNames': ['真歩流'], 'conflict': 'skip'}
    r = client.post(f'/api/projects/{src["id"]}/casts/copy', json=req)
    assert r.status_code == 200
    assert r.json()['summary']['copied'] == 0 and r.json()['summary']['skipped'] == ['真歩流']
    assert r.json()['project']['revision'] == tgt['revision']
    req['castNames'] = ['真歩流', '愛理']
    req['conflict'] = 'rename'
    r = client.post(f'/api/projects/{src["id"]}/casts/copy', json=req)
    assert r.status_code == 200, r.text
    voices = r.json()['project']['voices']
    assert [v['name'] for v in voices] == ['真歩流', '愛理', '愛理（コピー）', '真歩流（コピー）', '愛理（コピー2）']
    assert voices[0]['description'] == '変更禁止'
    assert r.json()['summary']['copied'] == 2


def test_copy_casts_rejects_invalid_request_and_stale_revision_without_mutation(client):
    src = project(client)
    tgt = client.post('/api/projects', json={'title': '続編'}).json()
    src['voices'] = [{'name': '零度', 'refs': []}]
    src = client.put(f'/api/projects/{src["id"]}', json={'state': src, 'expectedRevision': src['revision']}).json()
    route = f'/api/projects/{src["id"]}/casts/copy'
    base = {'targetProjectId': tgt['id'], 'castNames': ['零度']}
    assert client.post(route, json={**base, 'targetProjectId': src['id']}).status_code == 400
    assert client.post(route, json={**base, 'castNames': ['零度', '零度']}).status_code == 400
    assert client.post(route, json={**base, 'castNames': ['存在しない人物']}).status_code == 400
    assert client.post(route, json={**base, 'expectedSourceRevision': 0}).status_code == 409
    assert client.post(route, json={**base, 'expectedTargetRevision': 0}).status_code == 409
    assert client.get(f'/api/projects/{tgt["id"]}').json()['revision'] == tgt['revision']
    assert not list((core.workdir(tgt['id']) / 'assets').iterdir())


def test_copy_casts_missing_source_reference_is_atomic(client):
    src = project(client)
    tgt = client.post('/api/projects', json={'title': '続編'}).json()
    ref = upload(client, src['id'], kind='reference')
    src = client.get(f'/api/projects/{src["id"]}').json()
    src['voices'] = [{'name': '零度', 'refs': [ref['id']]}]
    src = client.put(f'/api/projects/{src["id"]}', json={'state': src, 'expectedRevision': src['revision']}).json()
    (core.asset_path(src['id'], ref)).unlink()
    r = client.post(f'/api/projects/{src["id"]}/casts/copy', json={
        'targetProjectId': tgt['id'], 'castNames': ['零度']})
    assert r.status_code == 400, r.text
    assert client.get(f'/api/projects/{tgt["id"]}').json()['revision'] == tgt['revision']
    assert not list((core.workdir(tgt['id']) / 'assets').iterdir())


def test_cast_copy_ui_has_bulk_single_and_all_actions():
    js = (Path(__file__).resolve().parents[1] / 'static' / 'app.js').read_text(encoding='utf-8')
    assert 'openCastCopyBtn' in js
    assert 'data-cast-copy' in js
    assert 'castCopyAllBtn' in js and 'castCopyNoneBtn' in js
    assert 'castCopyConfirmBtn' in js
    assert '/casts/copy' in js

# v0.1.7: Excel full-replace safety and backward compatibility

def _workbook_bytes(rows, *, sheet='第一幕', title_column=True):
    wb = openpyxl.Workbook(); ws=wb.active; ws.title=sheet
    ws.append(['台詞ID','シーン','話者','台詞','キャプション','間(秒)'] if title_column else ['備考'])
    for row in rows: ws.append(row)
    out=io.BytesIO();wb.save(out);return out.getvalue()


def test_excel_replace_discards_old_script_but_keeps_audio_and_casts(client):
    s=project(client); pid=s['id']; sid=s['scenes'][0]['id']
    s['scenes'][0]['name']='第一幕'; discarded_scene=core.uid('scene'); s['scenes'].append({'id':discarded_scene,'name':'第二幕'})
    retained=core.new_line(sid,'真歩流','残す台詞'); removed=core.new_line(sid,'愛理','消す台詞')
    removed2=core.new_line(discarded_scene,'零度','削除シーンの台詞')
    s['lines']=[retained,removed,removed2];s['voices']=[{'name':'真歩流','refs':[],'description':'探偵'}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    a=upload(client,pid,kind='voice');ref=upload(client,pid,kind='reference')
    s=client.get(f'/api/projects/{pid}').json();s['voices'][0]['refs']=[ref['id']]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    s=client.post(f'/api/projects/{pid}/lines/{retained["id"]}/attach/{a["id"]}').json()
    s['clips'] += [
        {'id':'clip_fx_old','sceneId':sid,'trackId':'sfx','assetId':a['id'],'start':.2,'duration':.2,'anchorLineId':removed['id']},
        {'id':'clip_fx_independent','sceneId':sid,'trackId':'sfx','assetId':a['id'],'start':.1,'duration':.2},
        {'id':'clip_fx_old_scene','sceneId':discarded_scene,'trackId':'sfx','assetId':a['id'],'start':.1,'duration':.2},
    ]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    orig=s.copy()
    data=_workbook_bytes([[retained['id'],'第一幕','真歩流','残す台詞','驚いて',.6],
                          [None,'第一幕','愛理','新しい台詞','普通',.3]])
    r=client.post(f'/api/projects/{pid}/excel/import',data={'mode':'replace','expectedRevision':s['revision']},
       files={'file':('correct.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert r.status_code==200,r.text
    result=r.json();s=result['project'];summary=result['summary']
    assert summary['mode']=='replace' and summary['removed']==2 and summary['removedScenes']==1
    assert len(s['lines'])==2 and s['lines'][0]['id']==retained['id']
    assert removed['id'] not in [x['id'] for x in s['lines']]
    assert all(x['name']!='第二幕' for x in s['scenes'])
    assert len(s['lines'][0]['takes'])==1 and s['lines'][0]['selectedTakeId']==orig['lines'][0]['selectedTakeId']
    assert {x['id'] for x in s['clips']}=={'clip_fx_independent',*{c['id'] for c in orig['clips'] if c.get('lineId')==retained['id']}}
    assert s['voices'][0]['description']=='探偵' and s['voices'][0]['refs']==[ref['id']]
    assert len(s['assets'])==2
    assert (core.asset_path(pid, ref)).exists()
    backup=core.workdir(pid)/'backups'/summary['backup']
    assert backup.exists()
    import json
    prior=json.loads(backup.read_text(encoding='utf-8'))
    assert prior['revision']==orig['revision'] and len(prior['lines'])==3
    assert prior['title']==orig['title']


def test_excel_merge_retains_unmentioned_lines(client):
    s=project(client);sid=s['scenes'][0]['id'];s['scenes'][0]['name']='第一幕'
    s['lines']=[core.new_line(sid,'真歩流','前の台詞')]
    s=client.put(f'/api/projects/{s["id"]}',json={'state':s,'expectedRevision':s['revision']}).json()
    x=_workbook_bytes([[None,'第一幕','愛理','新しい台詞',None,.5]])
    r=client.post(f'/api/projects/{s["id"]}/excel/import',files={'file':('new.xlsx',x)},data={'mode':'merge'})
    assert r.status_code==200,r.text
    assert len(r.json()['project']['lines'])==2 and r.json()['summary']['removed']==0
    assert not (core.workdir(s['id'])/'backups').exists()


def test_excel_replace_rejects_missing_data_duplicate_and_conflict(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['lines']=[core.new_line(sid,'真歩流','残る')]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    def attempt(data, **fields):
        return client.post(f'/api/projects/{pid}/excel/import',data={'mode':'replace','expectedRevision':s['revision'],**fields},files={'file':('bad.xlsx',data)})
    assert attempt(_workbook_bytes([],title_column=False)).status_code==400
    assert attempt(_workbook_bytes([])).status_code==400
    id=s['lines'][0]['id'];dupes=[[id,'シーン1','真歩流','1',None,.5],[id,'シーン1','真歩流','2',None,.5]]
    assert attempt(_workbook_bytes(dupes)).status_code==400
    good=_workbook_bytes([[None,'シーン1','真歩流','2',None,.5]])
    assert attempt(good,expectedRevision=0).status_code==409
    assert attempt(good,mode='invalid').status_code==400
    unchanged=client.get(f'/api/projects/{pid}').json()
    assert unchanged['revision']==s['revision'] and unchanged['lines'][0]['text']=='残る'
    assert not (core.workdir(pid)/'backups').exists()


def test_excel_ui_exposes_replace_and_confirmation():
    js=(Path(__file__).resolve().parents[1]/'static'/'app.js').read_text(encoding='utf8')
    assert 'chooseExcelImport' in js and "value=\"replace\"" in js
    assert 'if(mode===\'replace\' && !confirm(' in js
    assert "form.append('mode',mode)" in js



def test_cast_anime_model_is_default_for_legacy_and_new_lines(client):
    initial = project(client)
    pid = initial['id']
    current = client.get(f'/api/projects/{pid}').json()
    sid = current['scenes'][0]['id']
    anime = 'phasefield-audio/Irodori-TTS-v4.1-Anime'
    current['voices'] = [{'name': '真歩流', 'engine': 'irodori', 'mode': 'clone', 'model': anime, 'refs': []}]
    # Old lines already stored Small; newer lines explicitly track cast defaults.
    for i in range(164):
        item = core.new_line(sid, '真歩流', f'セリフ{i}')
        if i < 163:
            item.pop('ttsSource', None)
        current['lines'].append(item)
    saved = client.put(f'/api/projects/{pid}', json={
        'state': current, 'expectedRevision': current['revision']}).json()
    result = client.post(f'/api/projects/{pid}/generate/preview', json={
        'lineIds': [line['id'] for line in saved['lines']]})
    assert result.status_code == 200, result.text
    assert result.json()['total'] == 164
    assert result.json()['groups'] == [{'engine':'irodori','model':anime,'count':164}]
    assert core.resolve_line_tts(saved['lines'][0], saved['voices'][0])['model'] == anime
    assert core.resolve_line_tts(saved['lines'][-1], saved['voices'][0])['source'] == 'cast'


def test_individual_override_keeps_custom_small_even_with_anime_cast(client):
    s=project(client)
    anime = 'phasefield-audio/Irodori-TTS-v4.1-Anime'
    s['voices'] = [{'name':'愛理','engine':'irodori','model':anime,'mode':'clone','refs':[]}]
    sid = s['scenes'][0]['id']
    normal = core.new_line(sid,'愛理','通常はAnime')
    explicit = core.new_line(sid,'愛理','このセリフだけSmall')
    explicit['ttsSource'] = 'line'
    legacy_custom = core.new_line(sid,'愛理','旧版のカスタムモデル')
    legacy_custom.pop('ttsSource',None)
    legacy_custom['model'] = 'Aratako/Irodori-TTS-v4-Large'
    s['lines'] = [normal,explicit,legacy_custom]
    r=client.put(f'/api/projects/{s["id"]}',json={'state':s,'expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    preview=client.post(f'/api/projects/{s["id"]}/generate/preview',json={
        'lineIds':[x['id'] for x in s['lines']]})
    assert preview.status_code==200,preview.text
    assert preview.json()['groups']==[
        {'engine':'irodori','model':anime,'count':1},
        {'engine':'irodori','model':'Aratako/Irodori-TTS-v4.1-Small','count':1},
        {'engine':'irodori','model':'Aratako/Irodori-TTS-v4-Large','count':1},
    ]


def test_actual_generation_sends_cast_anime_model(client,monkeypatch):
    import app as api
    import json
    s=project(client)
    anime='phasefield-audio/Irodori-TTS-v4.1-Anime'
    s['voices']=[{'name':'真歩流','engine':'irodori','model':anime,'mode':'clone','refs':[]}]
    sid=s['scenes'][0]['id']
    line=core.new_line(sid,'真歩流','テスト用台詞')
    line.pop('ttsSource',None)
    s['lines']=[line]
    s=client.put(f'/api/projects/{s["id"]}',json={'state':s,'expectedRevision':s['revision']}).json()
    monkeypatch.setattr(api,'python_for',lambda *_:sys.executable)
    captured=[]
    def fake_run(cmd,**kw):
        job=json.loads(Path(cmd[-2]).read_text(encoding='utf-8'))
        captured.append(job)
        items=[]
        for it in job['items']:
            out=Path(it['output'])
            sf.write(out,np.zeros(4410,dtype='float32'),44100)
            items.append({'id':it['id'],'ok':True,'output':str(out)})
        Path(cmd[-1]).write_text(json.dumps({'ok':True,'items':items}),encoding='utf-8')
        return type('Result',(),{'returncode':0})()
    monkeypatch.setattr(api.subprocess,'run',fake_run)
    jid=core.uid('job')
    api.JOBS[jid]={'status':'running','completed':0,'total':0,'errors':[],'message':''}
    api.run_generation(jid,s['id'],[line['id']],40)
    assert api.JOBS[jid]['status']=='done',api.JOBS[jid]
    assert captured[0]['model']==anime
    done=client.get(f'/api/projects/{s["id"]}').json()
    assert len(done['lines'][0]['takes'])==1
    assert done['lines'][0]['status']=='ready'


def test_legacy_small_model_with_anime_cast_marks_previous_take_stale(client):
    s=project(client)
    sid=s['scenes'][0]['id']
    line=core.new_line(sid,'真歩流','こんにちは')
    line.pop('ttsSource',None)
    s['lines']=[line]
    s=client.put(f'/api/projects/{s["id"]}',json={'state':s,'expectedRevision':s['revision']}).json()
    audio=upload(client,s['id'],kind='voice')
    s=client.post(f'/api/projects/{s["id"]}/lines/{line["id"]}/attach/{audio["id"]}').json()
    assert s['lines'][0]['status']=='ready'
    s['voices']=[{'name':'真歩流','engine':'irodori','model':'phasefield-audio/Irodori-TTS-v4.1-Anime','refs':[]}]
    s=client.put(f'/api/projects/{s["id"]}',json={'state':s,'expectedRevision':s['revision']}).json()
    assert s['lines'][0]['status']=='stale'
    assert client.get(f'/api/projects/{s["id"]}').json()['lines'][0]['status']=='stale'


def test_tts_cast_defaults_ui_and_preview_are_available():
    js=(Path(__file__).resolve().parents[1]/'static'/'app.js').read_text(encoding='utf-8')
    assert "l.ttsSource='cast'" in js
    assert "v.model||'Aratako/Irodori-TTS-v4.1-Small'" in js
    assert '/generate/preview' in js
    assert 'lineTtsSource' in js
    assert 'この台詞だけ個別設定' in js
    assert 'このキャストの全台詞へ設定を適用' in js



def test_script_clear_replaces_wrong_import_and_retains_cast_audio(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['scenes'].append({'id':core.uid('scene'),'name':'間違った第二幕'})
    s['lines']=[core.new_line(sid,'真歩流','間違い'),core.new_line(s['scenes'][1]['id'],'愛理','間違い2')]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    ref=upload(client,pid,kind='reference');audio=upload(client,pid,kind='voice')
    s=client.get(f'/api/projects/{pid}').json()
    s['voices']=[{'name':'真歩流','refs':[ref['id']],'model':'phasefield-audio/Irodori-TTS-v4.1-Anime'}]
    s['clips']=[{'id':core.uid('clip'),'sceneId':sid,'trackId':'bgm','assetId':audio['id'],
                 'start':0,'duration':0.1}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    result=client.post(f'/api/projects/{pid}/script/clear',json={'expectedRevision':s['revision']})
    assert result.status_code==200,result.text
    data=result.json();cleared=data['project'];summary=data['summary']
    assert cleared['id']==pid and cleared['title']==s['title']
    assert cleared['lines']==[] and cleared['clips']==[]
    assert len(cleared['scenes'])==1 and cleared['scenes'][0]['name']=='シーン1'
    assert cleared['voices']==s['voices'] and cleared['assets']==s['assets']
    assert cleared['settings']==s['settings'] and cleared['tracks']==s['tracks']
    assert summary['removedLines']==2 and summary['removedClips']==1 and summary['removedScenes']==2
    backup=core.workdir(pid)/'backups'/summary['backup']
    assert backup.is_file()
    original=__import__('json').loads(backup.read_text(encoding='utf-8'))
    assert len(original['lines'])==2 and len(original['clips'])==1
    assert (core.asset_path(pid, audio)).is_file()
    assert (core.asset_path(pid, ref)).is_file()
    import io as _io
    import openpyxl as _openpyxl
    wb=_openpyxl.Workbook();ws=wb.active;ws.title='新しい第一幕'
    ws.append(['話者','台詞']);ws.append(['真歩流','正しい台詞'])
    buf=_io.BytesIO();wb.save(buf)
    response=client.post(f'/api/projects/{pid}/excel/import',files={'file':('new.xlsx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==200,response.text
    after=response.json()['project']
    assert len(after['lines'])==1 and after['lines'][0]['text']=='正しい台詞'
    assert [scene['name'] for scene in after['scenes']]==['新しい第一幕']
    assert after['voices'][0]['model']=='phasefield-audio/Irodori-TTS-v4.1-Anime'


def test_script_clear_rejects_conflicts_and_missing_project(client):
    current=project(client);pid=current['id']
    response=client.post(f'/api/projects/{pid}/script/clear',json={'expectedRevision':current['revision']-1})
    assert response.status_code==409
    assert client.get(f'/api/projects/{pid}').json()==current
    assert not (core.workdir(pid)/'backups').exists()
    missing=client.post('/api/projects/p_0000000000000000/script/clear',json={'expectedRevision':1})
    assert missing.status_code==404
    no_revision=client.post(f'/api/projects/{pid}/script/clear',json={})
    assert no_revision.status_code==422


def test_script_clear_frontend_has_explicit_confirmation():
    js=(Path(__file__).resolve().parents[1]/'static'/'app.js').read_text(encoding='utf-8')
    assert 'openClearScriptBtn' in js
    assert 'confirmClearScriptBtn' in js
    assert 'script/clear' in js
    assert 'removedLines' in js and 'removedClips' in js


def test_generated_audio_filename_defaults_custom_and_safe_names():
    l = {'id': 'line_1111111111111111'}
    assert core.generated_audio_filename(l, 'irodori') == 'line_1111111111111111 - irodori.wav'
    l['outputFilename'] = '真歩流_第01話'
    assert core.generated_audio_filename(l) == '真歩流_第01話.wav'
    l['outputFilename'] = '真歩流_第01話.WAV'
    assert core.generated_audio_filename(l) == '真歩流_第01話.wav'
    for bad in ('../bad', '..\\secret', 'bad/name', 'bad:name', '.hidden', 'CON',
                'COM1', 'hello.mp3', 'a..wav', 'bad\nfile', 'trailing.', 'x'*121):
        l['outputFilename'] = bad
        with pytest.raises(ValueError):
            core.generated_audio_filename(l)


def test_custom_filename_is_validated_at_save_and_does_not_invalidate_audio(client):
    s = project(client); pid=s['id']
    line=core.new_line(s['scenes'][0]['id'], '真歩流','こんにちは')
    s['lines']=[line]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    uploaded=upload(client,pid,kind='voice')
    s=client.post(f'/api/projects/{pid}/lines/{line["id"]}/attach/{uploaded["id"]}').json()
    current=s['lines'][0]
    assert current['status']=='ready'
    current['outputFilename']='take_new'
    saved=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert saved.status_code==200, saved.text
    s=saved.json()
    assert s['lines'][0]['status']=='ready'
    assert len(s['lines'][0]['takes'])==1
    assert s['assets'][0]['name']=='voice.wav'
    s['lines'][0]['outputFilename']='a/b.wav'
    bad=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']})
    assert bad.status_code==400


def test_custom_filename_generation_and_download(client,monkeypatch):
    import json
    import app as api
    from urllib.parse import unquote
    s=project(client); pid=s['id']
    one=core.new_line(s['scenes'][0]['id'],'真歩流','台詞1')
    two=core.new_line(s['scenes'][0]['id'],'真歩流','台詞2')
    one['outputFilename']='探偵001.wav'
    s['lines']=[one,two]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    monkeypatch.setattr(api,'python_for',lambda *_:sys.executable)
    def fake(cmd,**kwargs):
        data=json.loads(Path(cmd[-2]).read_text(encoding='utf-8'))
        rows=[]
        for item in data['items']:
            p=Path(item['output']);sf.write(p,np.zeros(4410,dtype='float32'),44100)
            rows.append({'id':item['id'],'ok':True,'output':str(p)})
        Path(cmd[-1]).write_text(json.dumps({'ok':True,'items':rows}),encoding='utf-8')
        return type('Ok',(),{'returncode':0})()
    monkeypatch.setattr(api.subprocess,'run',fake)
    jid=core.uid('job');api.JOBS[jid]={'status':'running','completed':0,'total':0,'errors':[],'message':''}
    api.run_generation(jid,pid,[one['id'],two['id']],40)
    assert api.JOBS[jid]['status']=='done',api.JOBS[jid]
    fresh=client.get(f'/api/projects/{pid}').json()
    by_id={x['id']:x for x in fresh['assets']}
    named=[]
    for line in fresh['lines']:
        aid=line['takes'][-1]['assetId']; named.append((line['id'],by_id[aid]['name']))
        response=client.get(f'/api/projects/{pid}/assets/{aid}/download')
        assert response.status_code==200
        assert len(response.content)>100
        filename=unquote(response.headers['content-disposition'])
        assert by_id[aid]['name'] in filename
        assert by_id[aid]['filename'].startswith(aid)  # Stored safely, distinct from display name
    assert named==[(one['id'],'探偵001.wav'),(two['id'],f'{two["id"]} - irodori.wav')]
    assert client.get(f'/api/projects/{pid}/assets/asset_missing/download').status_code==404


def test_excel_roundtrip_retains_optional_output_filename(client):
    s=project(client); pid=s['id']
    l=core.new_line(s['scenes'][0]['id'],'真歩流','事件です')
    l['outputFilename']='case_001'
    s['lines']=[l]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    xlsx=client.get(f'/api/projects/{pid}/excel/export')
    w=openpyxl.load_workbook(io.BytesIO(xlsx.content)); sheet=w.active
    assert sheet.cell(1,7).value=='音声ファイル名'
    assert sheet.cell(2,7).value=='case_001'
    sheet.cell(2,7).value='case_002.wav'
    output=io.BytesIO();w.save(output)
    imported=client.post(f'/api/projects/{pid}/excel/import',files={'file':('script.xlsx',output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert imported.status_code==200,imported.text
    line=imported.json()['project']['lines'][0]
    assert line['id']==l['id'] and line['outputFilename']=='case_002.wav'
    assert core.generated_audio_filename(line)=='case_002.wav'


def test_script_clear_and_filename_ui_have_controls():
    js=(Path(__file__).resolve().parents[1]/'static'/'app.js').read_text(encoding='utf-8')
    assert 'data-line-field="outputFilename"' in js
    assert 'id="downloadTakeBtn"' in js
    assert 'confirmClearScriptBtn' in js


def test_global_tts_settings_shared_across_projects(client, tmp_path):
    """PC-wide TTS paths survive project switches, while gap remains project scoped."""
    p1 = project(client)
    p2 = client.post('/api/projects', json={'title': '別の作品'}).json()
    initial = client.get('/api/settings/global').json()
    expected = {
        'irodori_dir': str(tmp_path / 'shared-irodori'),
        'irodori_python': str(tmp_path / 'python-irodori.exe'),
        'qwen_python': str(tmp_path / 'python-qwen.exe'),
    }
    updated = client.put('/api/settings/global', json={
        **expected, 'expectedRevision':initial['revision'],
    })
    assert updated.status_code == 200, updated.text
    saved = updated.json()
    assert saved['revision'] == initial['revision'] + 1
    assert all(saved[k] == expected[k] for k in expected)
    assert (tmp_path / 'global_settings.json').exists()

    # Both projects remain structurally intact; project-specific timing survives.
    one = client.get(f'/api/projects/{p1["id"]}').json()
    two = client.get(f'/api/projects/{p2["id"]}').json()
    one['settings']['defaultGap'] = 1.3
    one = client.put(f'/api/projects/{p1["id"]}', json={'state':one, 'expectedRevision':one['revision']}).json()
    assert client.get(f'/api/projects/{p2["id"]}').json()['settings']['defaultGap'] == .65
    assert client.get(f'/api/projects/{p1["id"]}').json()['settings']['defaultGap'] == 1.3
    assert client.get('/api/settings/global').json() == saved

    # New projects also use the same global settings.
    fresh = client.post('/api/projects', json={'title':'三本目'}).json()
    assert fresh['settings'] == {'defaultGap': .65}
    assert client.get('/api/settings/global').json() == saved


def test_global_settings_legacy_migration_once_without_mutating_projects(client, tmp_path):
    """Migrate project-specific TTS directories, then never re-import stale paths."""
    old = project(client)
    old['settings']['irodori_dir'] = 'E:\\Irodori-Custom'
    old['settings']['irodori_python'] = 'E:\\Irodori-Custom\\venv\\python.exe'
    old = client.put(f'/api/projects/{old["id"]}', json={
        'state':old, 'expectedRevision':old['revision'],
    }).json()
    newer = client.post('/api/projects', json={'title':'次の作品'}).json()
    newer['settings']['irodori_dir'] = 'D:\\Irodori-TTS' # legacy default, not user choice
    newer['settings']['qwen_python'] = 'F:\\qwen\\python.exe'
    newer = client.put(f'/api/projects/{newer["id"]}', json={
        'state':newer, 'expectedRevision':newer['revision'],
    }).json()

    settings = client.get('/api/settings/global').json()
    assert settings['irodori_dir'] == 'E:\\Irodori-Custom'
    assert settings['irodori_python'] == 'E:\\Irodori-Custom\\venv\\python.exe'
    assert settings['qwen_python'] == 'F:\\qwen\\python.exe'
    # The original stored projects stay unmodified.
    assert client.get(f'/api/projects/{old["id"]}').json()['revision'] == old['revision']
    assert client.get(f'/api/projects/{newer["id"]}').json()['revision'] == newer['revision']
    newer['settings']['irodori_dir'] = 'G:\\wrong-old-project-path'
    client.put(f'/api/projects/{newer["id"]}', json={
        'state':newer, 'expectedRevision':newer['revision'],
    })
    assert client.get('/api/settings/global').json() == settings


def test_global_settings_validate_and_reject_stale(client):
    initial = client.get('/api/settings/global').json()
    old = initial['revision']
    replacement = {'irodori_dir':'C:\\Irodori', 'irodori_python':'',
                   'qwen_python':'', 'expectedRevision':old}
    assert client.put('/api/settings/global', json=replacement).status_code == 200
    assert client.put('/api/settings/global', json=replacement).status_code == 409
    payload = {**replacement,'expectedRevision':old+1,'irodori_dir':'broken\npath'}
    assert client.put('/api/settings/global', json=payload).status_code == 400
    assert client.get('/api/settings/global').json()['irodori_dir'] == 'C:\\Irodori'


def test_irodori_generation_uses_global_settings_not_old_project_paths(client, tmp_path, monkeypatch):
    import app as api
    s = project(client)
    pid = s['id']
    s['settings']['irodori_dir'] = 'C:\\OLD-PROJECT-ONLY'
    s = client.put(f'/api/projects/{pid}', json={'state':s,'expectedRevision':s['revision']}).json()
    global_path = tmp_path / 'shared-irodori'
    global_path.mkdir()
    (global_path / 'infer.py').write_text('# test',encoding='utf-8')
    (global_path / '.venv' / 'Scripts').mkdir(parents=True)
    exe = global_path / '.venv' / 'Scripts' / 'python.exe'
    exe.write_text('mock', encoding='utf-8')
    before = client.get('/api/settings/global').json()
    saved = client.put('/api/settings/global', json={
        'irodori_dir':str(global_path), 'irodori_python':'', 'qwen_python':'',
        'expectedRevision':before['revision'],
    })
    assert saved.status_code == 200
    assert api.python_for('irodori',core.load_project(pid)) == str(exe)
    # Voice generation payload and Voice Design both read the shared path.
    assert core.load_project(pid)['settings']['irodori_dir'] == 'C:\\OLD-PROJECT-ONLY'


def test_global_settings_ui_load_and_save_api_wiring():
    js = (Path(__file__).resolve().parents[1] / 'static' / 'app.js').read_text(encoding='utf-8')
    modal_js = js[js.index('async function openSettings()'):js.index('async function playPreview(')]
    assert "req('/api/settings/global')" in modal_js
    assert "req('/api/settings/global',json('PUT',settings))" in modal_js
    assert 'state.settings={...state.settings,defaultGap:gap}' in modal_js
    assert 'state.settings||{}' in modal_js


def test_selected_scene_takes_zip_uses_only_selected_take_and_custom_names(client):
    from zipfile import ZipFile
    state = project(client)
    pid, scene_id = state['id'], state['scenes'][0]['id']
    other_scene = {'id': core.uid('scene'), 'name': '第二幕'}
    state['scenes'].append(other_scene)
    lines = [core.new_line(scene_id,'真歩流','第一台詞'),
             core.new_line(scene_id,'愛理','第二台詞'),
             core.new_line(scene_id,'零度','未生成の台詞'),
             core.new_line(other_scene['id'],'零度','別シーン')]
    lines[0]['outputFilename'] = '会話.wav'
    lines[1]['outputFilename'] = '会話.wav'
    state['lines'] = lines
    saved = client.put(f'/api/projects/{pid}',json={
        'state':state,'expectedRevision':state['revision']})
    assert saved.status_code == 200, saved.text
    old_take = upload(client,pid,110)
    current = client.post(f'/api/projects/{pid}/lines/{lines[0]["id"]}/attach/{old_take["id"]}').json()
    chosen = upload(client,pid,220)
    current = client.post(f'/api/projects/{pid}/lines/{lines[0]["id"]}/attach/{chosen["id"]}').json()
    second = upload(client,pid,330)
    current = client.post(f'/api/projects/{pid}/lines/{lines[1]["id"]}/attach/{second["id"]}').json()
    separate = upload(client,pid,440)
    current = client.post(f'/api/projects/{pid}/lines/{lines[3]["id"]}/attach/{separate["id"]}').json()
    revision = current['revision']
    # Non-dialogue-only audio never joins a ZIP of selected line takes.
    extra = upload(client,pid,555,kind='bgm')
    r=client.get(f'/api/projects/{pid}/scenes/{scene_id}/selected-takes.zip')
    assert r.status_code == 200,r.text
    assert r.headers['content-type'] == 'application/zip'
    assert r.headers['x-vds-selected-takes'] == '2'
    assert 'attachment' in r.headers['content-disposition']
    with ZipFile(io.BytesIO(r.content)) as archive:
        assert archive.namelist() == ['会話.wav','会話 (2).wav']
        assert archive.read('会話.wav') == (core.asset_path(pid, chosen)).read_bytes()
        assert archive.read('会話 (2).wav') == (core.asset_path(pid, second)).read_bytes()
    after=client.get(f'/api/projects/{pid}').json()
    assert after['revision'] == revision+1  # one extra asset upload, no mutation by ZIP
    assert after['lines'][0]['takes'][0]['assetId'] == old_take['id']
    assert len(after['assets']) == 5


def test_selected_scene_zip_uses_asset_names_when_no_custom_filename_and_avoids_case_collisions(client):
    from zipfile import ZipFile
    state=project(client)
    pid, scene_id = state['id'], state['scenes'][0]['id']
    a = core.new_line(scene_id,'A','一')
    b = core.new_line(scene_id,'B','二')
    c = core.new_line(scene_id,'C','三')
    state['lines']=[a,b,c]
    state=client.put(f'/api/projects/{pid}',json={'state':state,'expectedRevision':state['revision']}).json()
    audio=upload(client,pid)
    for line in (a,b,c):
        state=client.post(f'/api/projects/{pid}/lines/{line["id"]}/attach/{audio["id"]}').json()
    # Names are inherited from asset metadata when not explicitly set on the line.
    state['assets'][0]['name']='VOICE.wav'
    state['lines'][1]['outputFilename']='voice.wav'
    state['lines'][2]['outputFilename']='voIce.wav'
    state=client.put(f'/api/projects/{pid}',json={'state':state,'expectedRevision':state['revision']}).json()
    r=client.get(f'/api/projects/{pid}/scenes/{scene_id}/selected-takes.zip')
    assert r.status_code==200,r.text
    with ZipFile(io.BytesIO(r.content)) as archive:
        assert archive.namelist()==['VOICE.wav','voice (2).wav','voIce (3).wav']
        assert all(archive.read(name)==(core.asset_path(pid, audio)).read_bytes()
                   for name in archive.namelist())


def test_selected_scene_zip_rejects_empty_invalid_scene_and_missing_audio(client):
    state=project(client)
    pid, scene_id=state['id'],state['scenes'][0]['id']
    route=f'/api/projects/{pid}/scenes/{scene_id}/selected-takes.zip'
    assert client.get(route).status_code == 400
    assert client.get(f'/api/projects/{pid}/scenes/scene_unknown/selected-takes.zip').status_code == 400
    line=core.new_line(scene_id,'A','こんにちは')
    state['lines']=[line]
    state=client.put(f'/api/projects/{pid}',json={'state':state,'expectedRevision':state['revision']}).json()
    asset=upload(client,pid)
    client.post(f'/api/projects/{pid}/lines/{line["id"]}/attach/{asset["id"]}')
    (core.asset_path(pid, asset)).unlink()
    r=client.get(route)
    assert r.status_code==404


def test_scene_zip_ui_button_wiring():
    js=(Path(__file__).resolve().parents[1]/'static'/'app.js').read_text(encoding='utf-8')
    assert 'downloadSceneTakesZipBtn' in js
    assert 'downloadSceneSelectedTakes()' in js
    assert '/selected-takes.zip' in js
    assert "await saveNow()" in js


def test_resizable_layout_contains_three_accessible_splitters():
    root=Path(__file__).resolve().parents[1]
    html=(root/'static'/'index.html').read_text(encoding='utf-8')
    for handle in ('resizeSidebar','resizeEditor','resizeInspector'):
        assert html.count(f'id="{handle}"')==1
    assert html.count('role="separator"')==3
    assert 'aria-orientation="horizontal"' in html
    assert 'aria-orientation="vertical"' in html
    assert '/static/app.js?v=0313' in html
    assert '/static/style.css?v=0313' in html


def test_resizable_layout_saves_and_restores_without_mutating_project():
    root=Path(__file__).resolve().parents[1]
    js=(root/'static'/'app.js').read_text(encoding='utf-8')
    css=(root/'static'/'style.css').read_text(encoding='utf-8')
    assert 'installPanelResizers();' in js
    assert "window.localStorage.setItem(PANEL_SIZE_KEY" in js
    assert 'window.localStorage.getItem(PANEL_SIZE_KEY)' in js
    assert "window.localStorage.removeItem(PANEL_SIZE_KEY)" in js
    assert "grip.addEventListener('pointerdown'" in js
    assert "grip.addEventListener('pointermove'" in js
    assert "grip.addEventListener('dblclick',resetPanelSizes)" in js
    assert "grip.addEventListener('keydown'" in js
    assert "--vds-editor-width" in css
    assert "--vds-sidebar-width" in css
    assert "--vds-inspector-height" in css
    assert '@media(max-width:950px)' in css
    assert '.panel-resizer-vertical{display:none}' in css


def test_delete_unused_cast_preserves_audio_and_other_casts(client):
    s=project(client);pid=s['id']
    reference=upload(client,pid,330,'reference')
    spoken=upload(client,pid,440,'voice')
    s=client.get(f'/api/projects/{pid}').json()
    s['voices']=[
        {'name':'間違えて作ったキャスト','engine':'irodori','model':'phasefield-audio/Irodori-TTS-v4.1-Anime',
         'mode':'clone','description':'','language':'Japanese','refText':'','refs':[reference['id']]},
        {'name':'正しいキャスト','engine':'irodori','model':'Aratako/Irodori-TTS-v4.1-Small',
         'mode':'clone','description':'','language':'Japanese','refText':'','refs':[]},
    ]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    original_revision=s['revision']
    r=client.post(f'/api/projects/{pid}/casts/delete',json={
        'name':'間違えて作ったキャスト','expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    after=r.json()
    assert [v['name'] for v in after['voices']]==['正しいキャスト']
    assert after['revision']==original_revision+1
    assert after['assets']==s['assets']
    assert after['lines']==s['lines'] and after['clips']==s['clips']
    assert (core.asset_path(pid, reference)).exists()
    assert (core.asset_path(pid, spoken)).exists()
    assert client.get(f'/api/projects/{pid}').json()['voices']==after['voices']


def test_delete_cast_in_use_rejected_without_mutation(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['voices']=[{'name':'愛理','refs':[],'engine':'irodori','model':'Aratako/Irodori-TTS-v4.1-Small'}]
    s['lines'].append(core.new_line(sid,'愛理','まだ使用中です'))
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    r=client.post(f'/api/projects/{pid}/casts/delete',json={'name':'愛理','expectedRevision':s['revision']})
    assert r.status_code==409
    assert '1件の台詞' in r.json()['detail']
    assert client.get(f'/api/projects/{pid}').json()==s
    # Once the speaker is corrected, deletion becomes possible.
    s['lines'][0]['speaker']='真歩流'
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    r=client.post(f'/api/projects/{pid}/casts/delete',json={'name':'愛理','expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    assert r.json()['lines'][0]['speaker']=='真歩流'


def test_delete_cast_rejects_stale_missing_and_invalid(client):
    s=project(client);pid=s['id']
    s['voices']=[{'name':'試験','refs':[]}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    endpoint=f'/api/projects/{pid}/casts/delete'
    assert client.post(endpoint,json={'name':'試験','expectedRevision':s['revision']-1}).status_code==409
    assert client.post(endpoint,json={'name':'未登録','expectedRevision':s['revision']}).status_code==404
    assert client.post(endpoint,json={'name':'','expectedRevision':s['revision']}).status_code==400
    assert client.post(endpoint,json={'name':'試験'}).status_code==422
    assert len(client.get(f'/api/projects/{pid}').json()['voices'])==1


def test_cast_delete_ui_wired_and_cache_busted():
    root=Path(__file__).resolve().parents[1]
    js=(root/'static/app.js').read_text(encoding='utf-8')
    html=(root/'static/index.html').read_text(encoding='utf-8')
    assert 'data-cast-delete' in js
    assert 'openCastDelete(Number(castDelete.dataset.castDelete))' in js
    assert 'confirmCastDeleteBtn' in js
    assert '/casts/delete' in js
    assert 'await saveNow();' in js[js.index('async function deleteCast'):js.index('function applyVoice(index)')]
    assert '/static/app.js?v=0313' in html


def test_asset_files_are_separated_and_servable(client):
    s=project(client);pid=s['id']
    ref=upload(client,pid,180,'reference')
    voice=upload(client,pid,220,'voice')
    sfx=upload(client,pid,330,'sfx')
    assert core.asset_path(pid,ref).parent.name=='reference_audio'
    assert core.asset_path(pid,voice).parent.name=='dialogue_audio'
    assert core.asset_path(pid,sfx).parent.name=='assets'
    for asset in (ref,voice,sfx):
        assert core.asset_path(pid,asset).is_file()
        response=client.get(f'/api/projects/{pid}/assets/{asset["id"]}')
        assert response.status_code==200,response.text
        assert response.content==core.asset_path(pid,asset).read_bytes()
        downloaded=client.get(f'/api/projects/{pid}/assets/{asset["id"]}/download')
        assert downloaded.status_code==200
        assert downloaded.content==response.content
    assert not list((core.workdir(pid)/'assets').glob('asset_*.wav')) or len(list((core.workdir(pid)/'assets').glob('asset_*.wav')))==1


def test_existing_project_migrates_audio_without_changing_ids_takes_or_casts(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    l=core.new_line(sid,'愛理','確認です')
    s['lines'].append(l)
    s['voices']=[{'name':'愛理','engine':'irodori','model':'demo',
                  'mode':'clone','language':'Japanese','refs':[]}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    ref=upload(client,pid,180,'reference')
    dialogue=upload(client,pid,220,'voice')
    effect=upload(client,pid,330,'sfx')
    s=client.post(f'/api/projects/{pid}/lines/{l["id"]}/attach/{dialogue["id"]}').json()
    s['voices'][0]['refs']=[ref['id']]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    revision=s['revision']
    original={a['id']:core.asset_path(pid,a).read_bytes() for a in (ref,dialogue,effect)}
    legacy=core.workdir(pid)/'assets'
    for asset in (ref,dialogue):
        core.asset_path(pid,asset).rename(legacy/asset['filename'])
    assert not core.asset_path(pid,ref).exists()
    assert not core.asset_path(pid,dialogue).exists()
    # Opening the project triggers migration from the pre-v0.2.5 flat layout.
    restored=client.get(f'/api/projects/{pid}')
    assert restored.status_code==200,restored.text
    newer=restored.json()
    assert newer['revision']==revision
    assert newer['voices'][0]['refs']==[ref['id']]
    assert newer['lines'][0]['selectedTakeId']==s['lines'][0]['selectedTakeId']
    for asset in (ref,dialogue,effect):
        assert core.asset_path(pid,asset).read_bytes()==original[asset['id']]
    assert not (legacy/ref['filename']).exists()
    assert not (legacy/dialogue['filename']).exists()
    assert (legacy/effect['filename']).is_file()
    zip_url=f'/api/projects/{pid}/scenes/{sid}/selected-takes.zip'
    assert client.get(zip_url).status_code==200
    assert client.get(f'/api/projects/{pid}/assets/{ref["id"]}').status_code==200


def test_reference_audio_trash_and_restore_in_separate_folder(client):
    s=project(client);pid=s['id']
    ref=upload(client,pid,200,'reference')
    location=core.asset_path(pid,ref)
    assert location.exists()
    deleted=client.delete(f'/api/projects/{pid}/assets/{ref["id"]}')
    assert deleted.status_code==200,deleted.text
    assert not location.exists()
    assert (core.workdir(pid)/'trash'/ref['filename']).exists()
    result=client.post(f'/api/projects/{pid}/trash/assets/{ref["id"]}/restore')
    assert result.status_code==200,result.text
    assert location.exists()
    assert client.get(f'/api/projects/{pid}/assets/{ref["id"]}').status_code==200


def test_voice_design_generated_reference_has_preview_id_and_correct_location(client,monkeypatch):
    import app as api
    import json
    s=project(client);pid=s['id']
    s['voices']=[{'name':'愛理','engine':'qwen','model':'Qwen/Qwen3-TTS-12Hz-1.7B-VoiceDesign',
                'mode':'design','description':'さわやかな声','language':'Japanese','refText':'','refs':[]}]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    monkeypatch.setattr(api,'python_for',lambda *_:sys.executable)
    def fake_run(cmd,**kwargs):
        job=json.loads(Path(cmd[-2]).read_text(encoding='utf8'))
        source=Path(job['items'][0]['output'])
        sf.write(source,np.zeros(4410,dtype='float32'),44100)
        Path(cmd[-1]).write_text(json.dumps({'ok':True,'items':[{'id':'愛理','ok':True,'output':str(source)}]}))
        return type('Completed',(),{'returncode':0})()
    monkeypatch.setattr(api.subprocess,'run',fake_run)
    jid=core.uid('job');api.JOBS[jid]={'status':'running','total':1,'completed':0,'errors':[],'message':''}
    api.run_voice_design(jid,pid,'愛理','こんにちは')
    assert api.JOBS[jid]['status']=='done'
    aid=api.JOBS[jid]['referenceAssetId']
    state=client.get(f'/api/projects/{pid}').json()
    asset=next(a for a in state['assets'] if a['id']==aid)
    assert aid in state['voices'][0]['refs']
    assert core.asset_path(pid,asset).parent.name=='reference_audio'
    assert client.get(f'/api/projects/{pid}/assets/{aid}').status_code==200
    js=(Path(__file__).resolve().parents[1]/'static'/'app.js').read_text('utf8')
    assert 'cast-reference-preview' in js
    assert 'previewDesignedReferenceBtn' in js
    assert 'data-preview=' in js


def test_cast_audio_preview_and_controls_stay_within_one_card():
    root = Path(__file__).resolve().parents[1]
    js = (root / 'static' / 'app.js').read_text(encoding='utf-8')
    css = (root / 'static' / 'style.css').read_text(encoding='utf-8')
    assert '<article class="cast-card" data-voice="${i}"' in js
    assert '<header class="cast-card-heading">' in js
    assert '<section class="cast-reference-section">' in js
    assert '${esc(v.name)} の参照音声' in js
    assert '<div class="asset-actions cast-card-actions">' in js
    assert 'data-preview="${esc(a.id)}"' in js
    assert 'data-cast-delete="${i}"' in js
    assert '</div></article>`).join' in js
    assert '.cast-card-actions' in css
    assert '.cast-reference-section' in css


def test_delete_scene_preserves_audio_casts_and_other_scenes(client):
    import json
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    extra=core.uid('scene');s['scenes'].append({'id':extra,'name':'食堂'})
    s['voices']=[{'name':'真歩流','engine':'irodori','model':core.DEFAULT_IRODORI_MODEL,'mode':'clone','refs':[]}]
    bad=core.new_line(sid,'真歩流','消えるシーン')
    good=core.new_line(extra,'真歩流','残るシーン')
    s['lines']=[bad,good]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    sound=upload(client,pid,kind='voice')
    s=client.get(f'/api/projects/{pid}').json()
    core.attach_take(s,s['lines'][0],sound)
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    aid=sound['id'];location=core.asset_path(pid,sound)
    assert location.exists()
    bad_line_id=bad['id']
    r=client.post(f'/api/projects/{pid}/scenes/{sid}/delete',json={'expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    body=r.json();updated=body['project']
    assert [x['id'] for x in updated['scenes']]==[extra]
    assert [x['id'] for x in updated['lines']]==[good['id']]
    assert all(c.get('sceneId')!=sid and c.get('lineId')!=bad_line_id for c in updated['clips'])
    assert updated['voices']==s['voices']
    assert [a['id'] for a in updated['assets']]==[aid]
    assert location.exists()
    backup=core.workdir(pid)/'backups'/body['summary']['backup']
    assert backup.is_file()
    assert json.loads(backup.read_text(encoding='utf-8'))['revision']==s['revision']
    assert body['summary']['removedLines']==1
    assert body['summary']['removedClips']==1


def test_delete_last_scene_creates_blank_scene(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    s['lines'].append(core.new_line(sid,'愛理','削除'))
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    result=client.post(f'/api/projects/{pid}/scenes/{sid}/delete',json={'expectedRevision':s['revision']})
    assert result.status_code==200,result.text
    after=result.json()['project']
    assert after['scenes']==[{'id':after['scenes'][0]['id'],'name':'シーン1'}]
    assert after['scenes'][0]['id']!=sid
    assert after['lines']==[] and after['clips']==[]


def test_scene_delete_revision_conflict_and_missing_are_atomic(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    original=(core.workdir(pid)/'project.json').read_bytes()
    assert client.post(f'/api/projects/{pid}/scenes/{sid}/delete',json={'expectedRevision':s['revision']-1}).status_code==409
    assert client.post(f'/api/projects/{pid}/scenes/no-such-scene/delete',json={'expectedRevision':s['revision']}).status_code==404
    assert (core.workdir(pid)/'project.json').read_bytes()==original
    assert not (core.workdir(pid)/'backups').exists()


def test_delete_selected_take_moves_media_to_trash_and_falls_back(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    line=core.new_line(sid,'真歩流','サブレ');s['lines']=[line]
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    old=upload(client,pid,freq=440,kind='voice');new=upload(client,pid,freq=550,kind='voice')
    s=client.get(f'/api/projects/{pid}').json()
    core.attach_take(s,s['lines'][0],old)
    core.attach_take(s,s['lines'][0],new)
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    line=s['lines'][0];selected=line['selectedTakeId'];previous=line['takes'][0]['id']
    r=client.post(f'/api/projects/{pid}/lines/{line["id"]}/takes/{selected}/delete',json={'expectedRevision':s['revision']})
    assert r.status_code==200,r.text
    after=r.json()['project'];l=after['lines'][0]
    assert len(l['takes'])==1 and l['selectedTakeId']==previous
    assert l['takes'][0]['assetId']==old['id']
    assert new['id'] not in {a['id'] for a in after['assets']}
    assert not core.asset_path(pid,new).exists()
    assert (core.workdir(pid)/'trash'/new['filename']).is_file()
    assert any(c['assetId']==old['id'] for c in after['clips'] if c.get('lineId')==l['id'])
    restored=client.post(f'/api/projects/{pid}/trash/assets/{new["id"]}/restore')
    assert restored.status_code==200,restored.text
    assert core.asset_path(pid,new).is_file()
    assert len(restored.json()['project']['lines'][0]['takes'])==1  # no implicit relink


def test_delete_take_protects_shared_media_and_revision(client):
    s=project(client);pid=s['id'];sid=s['scenes'][0]['id']
    a=core.new_line(sid,'愛理','その1');b=core.new_line(sid,'愛理','その2')
    s['lines']=[a,b];s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    audio=upload(client,pid,kind='voice');s=client.get(f'/api/projects/{pid}').json()
    core.attach_take(s,s['lines'][0],audio)
    core.attach_take(s,s['lines'][1],audio)
    s=client.put(f'/api/projects/{pid}',json={'state':s,'expectedRevision':s['revision']}).json()
    line=s['lines'][0];url=f'/api/projects/{pid}/lines/{line["id"]}/takes/{line["selectedTakeId"]}/delete'
    assert client.post(url,json={'expectedRevision':s['revision']-1}).status_code==409
    assert client.post(url,json={'expectedRevision':s['revision']}).status_code==409
    assert core.asset_path(pid,audio).exists()
    assert len(core.load_project(pid)['lines'][0]['takes'])==1
    assert client.post(f'/api/projects/{pid}/lines/{line["id"]}/takes/missing/delete',json={'expectedRevision':s['revision']}).status_code==404


def test_scene_and_take_delete_ui_wiring():
    root=Path(__file__).resolve().parents[1]
    js=(root/'static/app.js').read_text(encoding='utf-8')
    html=(root/'static/index.html').read_text(encoding='utf-8')
    assert 'data-delete-scene=' in js
    assert 'data-confirm-scene=' in js
    assert 'confirmSceneDelete(' in js
    assert '/scenes/${sid}/delete' in js
    assert 'id="deleteSelectedTakeBtn"' in js
    assert '/lines/${line.id}/takes/${take.id}/delete' in js
    assert '/static/app.js?v=0313' in html


def test_undo_redo_controls_are_present_and_cache_busted():
    root=Path(__file__).resolve().parents[1]
    html=(root/'static'/'index.html').read_text(encoding='utf-8')
    js=(root/'static'/'app.js').read_text(encoding='utf-8')
    assert 'id="undoBtn"' in html and 'id="redoBtn"' in html
    assert '/static/app.js?v=0313' in html
    assert '/static/style.css?v=0313' in html
    assert 'function undoEdit()' in js and 'function redoEdit()' in js
    assert "key==='z'&&!e.shiftKey" in js
    assert "key==='y'" in js
    assert "historyRestoring=true" in js
    assert "if(historyRestoring)return;handleBatchModalChange" in js
    assert "const HISTORY_LIMIT=80" in js
    assert "function clearEditHistory()" in js


def test_audio_library_cards_support_track_drop_and_time_position():
    root=Path(__file__).resolve().parents[1]
    js=(root/'static'/'app.js').read_text(encoding='utf-8')
    css=(root/'static'/'style.css').read_text(encoding='utf-8')
    assert 'draggable="true" data-drag-asset="${esc(a.id)}"' in js
    assert 'data-drop-track="${esc(track.id)}"' in js
    assert 'function onLibraryDragStart(e)' in js
    assert 'function onTimelineDrop(e)' in js
    assert 'function uploadToLane(files,trackId,atTime)' in js
    assert 'addClip(aid,track.id,seconds)' in js
    assert 'state=response.project;clearEditHistory();' in js
    assert "if(!track||track.kind==='dialogue')" in js
    assert '.tl-lane.timeline-drop-target' in css


def test_edit_history_never_rolls_back_asset_files():
    root=Path(__file__).resolve().parents[1]
    js=(root/'static'/'app.js').read_text(encoding='utf-8')
    snapshot=js[js.index('const HISTORY_FIELDS='):js.index('function recordEdit')]
    for key in ['title','scenes','lines','clips','tracks','voices','settings']:
        assert f"'{key}'" in snapshot
    assert "'assets'" not in snapshot
    assert 'setTimeout(()=>{historyRestoring=false},0)' in js
